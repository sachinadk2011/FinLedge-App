"""Safe one-time copying of legacy desktop workbooks into the shared data folder."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from .workbook_headers import header_key as _canonical_header

REPORT_NAME = ".finledge-data-migration.json"
WORKBOOKS = {
    "bank_transactions.xlsx": ("bank_transactions.xlsx", "bank.xlsx"),
    "share_transactions.xlsx": ("share_transactions.xlsx", "share.xlsx"),
    "personal_finance_bank_flow.xlsx": ("personal_finance_bank_flow.xlsx",),
    "personal_finance_cash_flow.xlsx": ("personal_finance_cash_flow.xlsx",),
    "personal_finance_transfer.xlsx": ("personal_finance_transfer.xlsx",),
}
REQUIRED_HEADER_GROUPS = {
    "bank_transactions.xlsx": (("date",), ("category", "particulars", "type"), ("amount", "total amount")),
    "share_transactions.xlsx": (("date",), ("share name", "company", "stock name"), ("category", "transaction type", "type")),
    "personal_finance_bank_flow.xlsx": (("date",), ("category",), ("amount", "total amount")),
    "personal_finance_cash_flow.xlsx": (("date",), ("category",), ("amount", "total amount")),
    "personal_finance_transfer.xlsx": (("date",), ("amount", "total amount"), ("direction", "flow type", "type")),
}
CANONICAL_SHEETS = {
    "bank_transactions.xlsx": "Bank",
    "share_transactions.xlsx": "Share",
    "personal_finance_bank_flow.xlsx": "Bank Flow",
    "personal_finance_cash_flow.xlsx": "Cash Flow",
    "personal_finance_transfer.xlsx": "Transfer Flow",
}
DERIVED_HEADERS = {
    "cumulative amount", "signed amount", "total amount", "profit/loss", "cumulative profit",
    "row id", "content hash", "created timestamp", "last updated timestamp", "updated device",
}


def _migration_timestamp() -> str:
    return datetime.now().astimezone().strftime("%Y-%m-%d %H%M")


def _unique_timestamped_directory(parent: Path, timestamp: str) -> Path:
    base = Path(parent) / timestamp
    candidate = base
    suffix = 2
    while candidate.exists():
        candidate = base.with_name(f"{base.name} ({suffix})")
        suffix += 1
    candidate.mkdir(parents=True)
    return candidate


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _validate_workbook(path: Path, canonical: str | None = None) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise ValueError("Workbook has no worksheets")
        expected = REQUIRED_HEADER_GROUPS.get(canonical, (("date",),))
        has_valid_header = False
        for sheet in workbook.worksheets:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            normalized = {_canonical_header(value) for value in first_row}
            if all(any(alias in normalized for alias in group) for group in expected):
                has_valid_header = True
            for _ in sheet.iter_rows(values_only=True):
                pass
        if not has_valid_header:
            raise ValueError("Workbook is missing required FinLedge columns")
    finally:
        workbook.close()


def _row_signatures(path: Path, fields_override: tuple[str, ...] | None = None) -> list[tuple[str, ...]]:
    from .workbook_merge import BUSINESS_FIELDS, _normalize_value, kind_for_workbook
    from .workbook_headers import header_key

    workbook = load_workbook(path, read_only=True, data_only=True)
    signatures: list[tuple[str, ...]] = []
    kind = kind_for_workbook(path.name)
    fields = fields_override or BUSINESS_FIELDS[kind]
    try:
        for sheet in workbook.worksheets:
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if not any(header_key(value) == "date" for value in first_row):
                continue
            headers = [header_key(value) for value in first_row]
            indexes = {header: index for index, header in enumerate(headers)}
            comparable_fields = [field for field in fields if header_key(field) in indexes]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(value is not None for value in row):
                    def value_for(field: str):
                        index = indexes.get(_canonical_header(field))
                        return row[index] if index is not None and index < len(row) else None

                signatures.append(tuple(
                        f"{field}={_normalize_value(field, value_for(field))}"
                        for field in comparable_fields
                    ))
    finally:
        workbook.close()
    return signatures


def _is_subset_workbook(source: Path, destination: Path) -> bool:
    try:
        from .workbook_merge import BUSINESS_FIELDS, kind_for_workbook
        from .workbook_headers import header_key

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            source_headers = set()
            for sheet in workbook.worksheets:
                source_headers.update(header_key(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        finally:
            workbook.close()
        kind = kind_for_workbook(source.name)
        comparable_fields = tuple(field for field in BUSINESS_FIELDS[kind] if header_key(field) in source_headers)
        source_counts = Counter(_row_signatures(source, comparable_fields))
        destination_counts = Counter(_row_signatures(destination, comparable_fields))
        return all(destination_counts[key] >= count for key, count in source_counts.items())
    except Exception:
        return False


def _copy_verified(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    if _sha256(source) != _sha256(destination):
        raise IOError(f"Could not verify copied workbook: {source.name}")


def _workbook_snapshot(path: Path, canonical: str) -> dict:
    _validate_workbook(path, canonical)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {}
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            first_row = next(rows, ())
            row_hash = hashlib.sha256()
            row_count = 0
            for row in rows:
                if not any(value is not None for value in row):
                    continue
                row_hash.update(json.dumps([_normalize_cell(value) for value in row], ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
                row_hash.update(b"\n")
                row_count += 1
            sheets[sheet.title] = {
                "headers": [_normalize_cell(value) for value in first_row],
                "recordCount": row_count,
                "rowsSha256": row_hash.hexdigest(),
            }
        return {
            "sha256": _sha256(path),
            "sheetNames": workbook.sheetnames,
            "sheets": sheets,
        }
    finally:
        workbook.close()


def _copy_verified_workbook(source: Path, destination: Path, canonical: str) -> dict:
    source_snapshot = _workbook_snapshot(source, canonical)
    _copy_verified(source, destination)
    backup_snapshot = _workbook_snapshot(destination, canonical)
    if source_snapshot != backup_snapshot:
        raise IOError(f"Workbook backup verification failed: {source.name}")
    return source_snapshot


def _normalize_cell(value) -> str:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return "" if value is None else str(value)


def _has_data_rows(path: Path) -> bool:
    """Return whether a valid workbook contains at least one non-empty data row."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                if not any(_canonical_header(value) == "date" for value in first_row):
                    continue
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if any(value is not None and str(value).strip() for value in row):
                        return True
            return False
        finally:
            workbook.close()
    except Exception:
        return False


def _ensure_canonical_sheet(path: Path, canonical: str) -> None:
    expected_title = CANONICAL_SHEETS[canonical]
    workbook = load_workbook(path)
    try:
        if expected_title in workbook.sheetnames:
            return
        expected = REQUIRED_HEADER_GROUPS[canonical]
        for sheet in workbook.worksheets:
            headers = {_canonical_header(cell.value) for cell in sheet[1]}
            if all(any(alias in headers for alias in group) for group in expected):
                if expected_title in workbook.sheetnames:
                    raise ValueError(f"Workbook already contains a {expected_title} sheet")
                sheet.title = expected_title
                workbook.save(path)
                return
    finally:
        workbook.close()


def _merge_workbooks(source: Path, destination: Path, canonical: str) -> None:
    """Merge a legacy source into a staging copy using stable IDs and tombstones."""
    from . import bank_service, personal_finance_service, share_service
    from .workbook_merge import merge_workbook_files

    configuration = {
        "bank_transactions.xlsx": ("Bank", bank_service.HEADERS, "bank"),
        "share_transactions.xlsx": ("Share", share_service.HEADERS, "share"),
        "personal_finance_bank_flow.xlsx": ("Bank Flow", personal_finance_service.HEADERS, "personal_finance"),
        "personal_finance_cash_flow.xlsx": ("Cash Flow", personal_finance_service.HEADERS, "personal_finance"),
        "personal_finance_transfer.xlsx": ("Transfer Flow", personal_finance_service.HEADERS, "transfer"),
    }
    sheet_name, headers, kind = configuration[canonical]
    merge_workbook_files(
        source,
        destination,
        sheet_name=sheet_name,
        headers=headers,
        kind=kind,
        identity_scope=f"{canonical.casefold()}#{sheet_name.casefold()}",
    )


def _row_signatures_from_sheet(sheet) -> set[tuple[str, ...]]:
    headers = [_canonical_header(cell.value) for cell in sheet[1]]
    business_headers = sorted(
        header for header in headers
        if header and header not in DERIVED_HEADERS
    )
    indexes = {header: index for index, header in enumerate(headers)}
    signatures = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in row):
            signatures.add(tuple(
                f"{header}={_normalize_cell(row[indexes[header]] if indexes[header] < len(row) else None)}"
                for header in business_headers
            ))
    return signatures


def _unique_migration_directory(backup_dir: Path, fallback_backup_dir: Path | None = None) -> Path:
    timestamp = _migration_timestamp()
    roots = [Path(backup_dir).expanduser().resolve()]
    if fallback_backup_dir is not None:
        fallback = Path(fallback_backup_dir).expanduser().resolve()
        if fallback not in roots:
            roots.append(fallback)

    last_error = None
    for root in roots:
        try:
            return _unique_timestamped_directory(root / "migration", timestamp)
        except OSError as exc:
            last_error = exc

    if last_error:
        raise last_error
    raise OSError("Could not create a migration backup directory")


def backup_root_workbooks_before_data_folder(data_dir: Path, backup_dir: Path) -> list[str]:
    """Back up root-level workbooks before creating a missing data child folder."""
    data_dir = Path(data_dir).expanduser().resolve()
    source_root = data_dir.parent
    sources = sorted({source for aliases in WORKBOOKS.values() for name in aliases
                      if (source := source_root / name).is_file()}, key=lambda source: source.name.casefold())
    if not sources:
        return []

    backup_dir = Path(backup_dir).expanduser().resolve()
    backup_root = backup_dir / "migration" / "pre-folder"
    try:
        batch_dir = _unique_timestamped_directory(backup_root, _migration_timestamp())
    except OSError:
        fallback_root = data_dir.parent / "backups" / "migration" / "pre-folder"
        batch_dir = _unique_timestamped_directory(fallback_root, _migration_timestamp())
    backed_up = []
    try:
        for source in sources:
            staging = batch_dir / f".{source.stem}.{uuid4().hex}.tmp.xlsx"
            destination = batch_dir / source.name
            try:
                canonical = _canonical_for_source(source)
                _copy_verified_workbook(source, staging, canonical)
                os.replace(staging, destination)
                _copy_verified_workbook(source, destination, canonical)
                backed_up.append(str(destination))
            finally:
                staging.unlink(missing_ok=True)
    except Exception:
        # Keep any completed backups for recovery and do not proceed to create the
        # destination data folder after a partial backup.
        raise
    return backed_up


def _direct_sources(legacy_dirs: list[Path]) -> list[tuple[int, Path, str, str]]:
    sources = []
    seen: set[Path] = set()
    for root_index, root_value in enumerate(legacy_dirs):
        root = Path(root_value).expanduser().resolve()
        candidate_roots = [root]
        data_child = (root / "FinLedge Data").resolve()
        if data_child.is_dir():
            candidate_roots.append(data_child)
        for candidate_root in candidate_roots:
            for canonical, aliases in WORKBOOKS.items():
                for filename in aliases:
                    source = (candidate_root / filename).resolve()
                    if source.is_file() and source not in seen:
                        sources.append((root_index, source, canonical, filename))
                        seen.add(source)
    return sources


def _source_fingerprint(sources: list[tuple[int, Path, str, str]]) -> dict[str, list[int | float]]:
    fingerprint = {}
    for _, source, _, _ in sources:
        stat = source.stat()
        fingerprint[str(source)] = [stat.st_size, stat.st_mtime]
    return dict(sorted(fingerprint.items(), key=lambda item: item[0].casefold()))


def _read_completed_report(path: Path, fingerprint: dict) -> dict | None:
    try:
        report = json.loads(path.read_text(encoding="utf-8"))
        if report.get("complete") is True and report.get("sourceFingerprint") == fingerprint:
            if all((path.parent / name).is_file() for name in WORKBOOKS):
                report["alreadyCompleted"] = True
                report["status"] = "existing_data"
                return report
    except (OSError, ValueError, TypeError):
        pass
    return None


def _backup_direct_sources(sources: list[tuple[int, Path, str, str]], migration_dir: Path) -> dict[Path, Path]:
    backups: dict[Path, Path] = {}
    for source_index, (_root_index, source, _canonical, _filename) in enumerate(sources, start=1):
        backup = migration_dir / "src" / str(source_index) / source.name
        _copy_verified_workbook(source, backup, _canonical_for_source(source, _canonical))
        backups[source] = backup
    return backups


def _canonical_for_source(source: Path, configured: str | None = None) -> str:
    name = source.name.casefold()
    for canonical, aliases in WORKBOOKS.items():
        if name in {alias.casefold() for alias in aliases}:
            return canonical
    return configured or source.name


def _write_json_atomic(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    staged = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        staged.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
        os.replace(staged, path)
    finally:
        staged.unlink(missing_ok=True)


def _copy_backup_history(
    data_dir: Path,
    legacy_dirs: list[Path],
    backup_dir: Path,
    target_backup_dir: Path | None = None,
) -> list[str]:
    sources = [Path(root).expanduser().resolve() / "backups" for root in legacy_dirs]
    sources.append(data_dir / "backups")
    copied: list[str] = []
    target_root = Path(target_backup_dir or backup_dir).expanduser().resolve()
    fallback_root = (data_dir.parent / "backups").resolve()
    for source_root in dict.fromkeys(sources):
        if not source_root.is_dir() or source_root.resolve() == target_root:
            continue
        for source in source_root.rglob("*"):
            if not source.is_file():
                continue
            relative = source.relative_to(source_root)
            target = target_root / relative
            if target.exists():
                if _sha256(target) == _sha256(source):
                    continue
                target = target.with_name(f"{target.stem} - legacy-{_sha256(source)[:8]}{target.suffix}")
                if target.exists():
                    continue
            try:
                _copy_verified(source, target)
            except OSError:
                if target_root == fallback_root:
                    raise
                target = fallback_root / relative
                if target.exists():
                    if _sha256(target) == _sha256(source):
                        copied.append(str(target))
                        continue
                    target = target.with_name(f"{target.stem} - legacy-{_sha256(source)[:8]}{target.suffix}")
                _copy_verified(source, target)
            copied.append(str(target))
    return copied


def _backup_candidates(
    canonical: str,
    legacy_dirs: list[Path],
    backup_dir: Path,
    excluded: Path,
    additional_backup_dirs: list[Path] | None = None,
) -> list[Path]:
    stems = {Path(name).stem.casefold() for name in WORKBOOKS[canonical]}
    filenames = {name.casefold() for name in WORKBOOKS[canonical]}
    roots = [Path(root).expanduser().resolve() / "backups" for root in legacy_dirs]
    roots.extend([Path(backup_dir).resolve(), Path(backup_dir).resolve() / ".." / "backups"])
    roots.extend(Path(root).expanduser().resolve() for root in additional_backup_dirs or [])
    candidates: set[Path] = set()
    for root in dict.fromkeys(path.resolve() for path in roots):
        if not root.is_dir():
            continue
        for path in root.rglob("*.xlsx"):
            resolved = path.resolve()
            if resolved == excluded or "corrupted" in {part.casefold() for part in resolved.parts}:
                continue
            if path.name.casefold() in filenames or path.parent.name.casefold() in stems:
                candidates.add(resolved)
    return sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)


def migrate_legacy_workbooks(
    data_dir: Path,
    legacy_dirs: list[Path],
    backup_dir: Path | None = None,
    *,
    cleanup_root_dirs: list[Path] | None = None,
    shared_sync_dir: Path | None = None,
) -> dict:
    """Validate, back up, recover, and safely centralize legacy workbook files."""
    data_dir = Path(data_dir).expanduser().resolve()
    backup_dir = Path(backup_dir or os.getenv("FINLEDGE_BACKUP_DIR") or data_dir / "backups").expanduser().resolve()
    backup_dir.mkdir(parents=True, exist_ok=True)
    legacy_dirs = list(dict.fromkeys(Path(root).expanduser().resolve() for root in legacy_dirs if Path(root).expanduser().resolve() != data_dir))
    cleanup_roots = list(dict.fromkeys(Path(root).expanduser().resolve() for root in cleanup_root_dirs or []))
    sources = [
        item for item in _direct_sources(legacy_dirs)
        if item[1].resolve() != (data_dir / item[2]).resolve()
    ]
    source_fingerprint = _source_fingerprint(sources)
    report_path = data_dir / REPORT_NAME
    completed_report = _read_completed_report(report_path, source_fingerprint)
    if completed_report is not None:
        return completed_report
    existing_targets = [data_dir / name for name in WORKBOOKS if (data_dir / name).exists()]
    needs_migration_snapshot = any(
        _is_valid(source, canonical)
        and (
            not (data_dir / canonical).is_file()
            or not _is_valid(data_dir / canonical, canonical)
            or not _has_data_rows(data_dir / canonical)
            or not _is_subset_workbook(source, data_dir / canonical)
        )
        for _, source, canonical, _ in sources
    ) or any(
        target.exists() and not _is_valid(target) for target in existing_targets
    ) or any(
        source.parent == root
        for _, source, _, _ in sources
        for root in cleanup_roots
    )
    # Some packaged AppX environments can create/check the Documents backup root
    # in Electron but cannot create its nested migration snapshot from the Python
    # sidecar. Keep a second recovery location beside the active Roaming folder.
    fallback_backup_dir = data_dir.parent / "backups"
    migration_dir = (
        _unique_migration_directory(backup_dir, fallback_backup_dir)
        if needs_migration_snapshot
        else None
    )

    effective_backup_dir = backup_dir
    if migration_dir is not None and not migration_dir.is_relative_to(backup_dir):
        effective_backup_dir = migration_dir.parent.parent
    result = {
        "completedAt": datetime.now().astimezone().isoformat(),
        "complete": True,
        "alreadyCompleted": False,
        "status": "existing_data" if existing_targets else "migrated_legacy_data",
        "destination": str(data_dir),
        "backupDirectory": str(effective_backup_dir),
        **({"configuredBackupDirectory": str(backup_dir)} if effective_backup_dir != backup_dir else {}),
        "migrationBackup": str(migration_dir) if migration_dir else None,
        "files": [],
        "removedLegacyFiles": [],
        "retainedLegacyFiles": [],
        "sourceFingerprint": source_fingerprint,
        "cleanupRoots": [str(root) for root in cleanup_roots],
        "sourceBackups": {},
    }

    try:
        # Root backups remain outside FinLedge Data. Older backups nested inside
        # FinLedge Data or legacy AppData roots are copied beside it for recovery.
        result["copiedBackupFiles"] = _copy_backup_history(
            data_dir,
            legacy_dirs,
            backup_dir,
            target_backup_dir=effective_backup_dir,
        )

        source_backups = _backup_direct_sources(sources, migration_dir) if migration_dir else {}
        result["sourceBackups"] = {str(source): str(backup) for source, backup in source_backups.items()}
        data_dir.mkdir(parents=True, exist_ok=True)
        if migration_dir:
            for target in existing_targets:
                if target.exists() and not _is_valid(target):
                    _copy_verified(target, migration_dir / "dest" / target.name)

        for canonical in WORKBOOKS:
            target = data_dir / canonical
            target_valid = target.is_file() and _is_valid(target, canonical)
            file_result = {"name": canonical, "destination": str(target)}
            direct = [source for _, source, name, _ in sources if name == canonical and _is_valid(source, canonical)]
            if target_valid and not _has_data_rows(target) and direct:
                if migration_dir:
                    _copy_verified(target, migration_dir / "dest" / target.name)
                selected = direct[0]
                staging = data_dir / f".migration-{Path(canonical).stem}-{uuid4().hex}.xlsx"
                try:
                    _copy_verified(selected, staging)
                    _validate_workbook(staging, canonical)
                    _ensure_canonical_sheet(staging, canonical)
                    _validate_workbook(staging, canonical)
                    os.replace(staging, target)
                finally:
                    staging.unlink(missing_ok=True)
                target_valid = _is_valid(target, canonical)
                file_result.update({"status": "replaced_empty_destination", "source": str(selected)})
            elif target_valid:
                file_result["status"] = "already_valid"
            else:
                damaged_existing = target.exists() and not _is_valid(target, canonical)
                damaged_sources_exist = any(
                    source_canonical == canonical and not _is_valid(source, canonical)
                    for _, source, source_canonical, _ in sources
                )
                backup_candidates = (
                    _backup_candidates(
                        canonical,
                        legacy_dirs,
                        backup_dir,
                        migration_dir or Path(),
                        additional_backup_dirs=[effective_backup_dir],
                    )
                    if damaged_existing or damaged_sources_exist else []
                )
                valid_backup = next((candidate for candidate in backup_candidates if _is_valid(candidate, canonical)), None)
                selected = direct[0] if direct else valid_backup
                if selected is None:
                    if target.exists():
                        raise RuntimeError(f"{canonical} is damaged and no valid backup or legacy copy was found. The original was left untouched.")
                    file_result["status"] = "not_present"
                    result["files"].append(file_result)
                    continue

                staging = data_dir / f".migration-{Path(canonical).stem}-{uuid4().hex}.xlsx"
                try:
                    _copy_verified(selected, staging)
                    _validate_workbook(staging, canonical)
                    _ensure_canonical_sheet(staging, canonical)
                    _validate_workbook(staging, canonical)
                    os.replace(staging, target)
                finally:
                    staging.unlink(missing_ok=True)
                if not _is_valid(target, canonical):
                    raise RuntimeError(f"Could not verify migrated workbook {canonical}.")
                target_valid = True
                file_result["status"] = "restored_from_backup" if selected in backup_candidates else "copied_from_legacy"
                file_result["source"] = str(selected)

            # A partially completed earlier run or a pre-existing destination can
            # coexist with newer root data. Back up the destination, merge unique
            # rows while preserving the source workbook in its original location.
            for _, source, source_canonical, _filename in sources:
                if source_canonical != canonical or not _is_valid(source, canonical) or not target_valid:
                    continue
                if _is_subset_workbook(source, target):
                    continue
                if migration_dir:
                    _copy_verified(target, migration_dir / "dest" / target.name)
                staging = data_dir / f".merge-{Path(canonical).stem}-{uuid4().hex}.xlsx"
                try:
                    _copy_verified(target, staging)
                    _merge_workbooks(source, staging, canonical)
                    _validate_workbook(staging, canonical)
                    os.replace(staging, target)
                finally:
                    staging.unlink(missing_ok=True)
                if not _is_valid(target, canonical) or not _is_subset_workbook(source, target):
                    raise RuntimeError(f"Could not verify merged workbook {canonical}; legacy data was retained.")
                file_result["status"] = "merged_legacy_data"

            if target_valid:
                file_result["validated"] = True
                for _, source, source_canonical, _filename in sources:
                    if source_canonical != canonical:
                        continue
                    result["retainedLegacyFiles"].append(str(source))
            result["files"].append(file_result)

        # Existing valid destination files were checked above. Missing modules are
        # initialized after this service returns. Legacy sources are never removed.
        if result["removedLegacyFiles"]:
            result["status"] = "migrated_legacy_data"
        elif result["retainedLegacyFiles"]:
            result["status"] = "source_data_retained_for_review"
        elif not existing_targets and not sources:
            result["status"] = "new_install"
        _write_json_atomic(report_path, result)
        return result
    except Exception as exc:
        result.update({"complete": False, "status": "failed", "error": str(exc)})
        try:
            _write_json_atomic(report_path, result)
        except OSError:
            pass
        raise RuntimeError(f"Could not safely prepare the FinLedge data folder: {exc}") from exc


def _source_identity_fingerprint(path: Path) -> dict:
    stat = path.stat()
    return {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns, "sha256": _sha256(path)}


def _metadata_pairs(path: Path) -> set[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    pairs = set()
    try:
        for sheet in workbook.worksheets:
            headers = [_canonical_header(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
            indexes = {header: index for index, header in enumerate(headers)}
            row_id_index, hash_index = indexes.get("row id"), indexes.get("content hash")
            if row_id_index is None or hash_index is None:
                continue
            for row in sheet.iter_rows(min_row=2, values_only=True):
                row_id = row[row_id_index] if row_id_index < len(row) else None
                content_hash = row[hash_index] if hash_index < len(row) else None
                if row_id and content_hash:
                    pairs.add((str(row_id), str(content_hash)))
    finally:
        workbook.close()
    return pairs


def _expected_metadata_pairs(source: Path, canonical: str) -> set[tuple[str, str]]:
    from . import bank_service, personal_finance_service, share_service
    from .workbook_merge import ensure_row_metadata, populate_missing_audit_metadata

    configuration = {
        "bank_transactions.xlsx": ("Bank", bank_service.HEADERS, "bank"),
        "share_transactions.xlsx": ("Share", share_service.HEADERS, "share"),
        "personal_finance_bank_flow.xlsx": ("Bank Flow", personal_finance_service.HEADERS, "personal_finance"),
        "personal_finance_cash_flow.xlsx": ("Cash Flow", personal_finance_service.HEADERS, "personal_finance"),
        "personal_finance_transfer.xlsx": ("Transfer Flow", personal_finance_service.HEADERS, "transfer"),
    }
    sheet_name, target_headers, kind = configuration[canonical]
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else next(
            candidate for candidate in workbook.worksheets
            if any(_canonical_header(value) == "date" for value in next(candidate.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        )
        source_headers = [str(value or "") for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        from .workbook_merge import map_row

        rows = [
            map_row(source_headers, row, target_headers)
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()
    rows = populate_missing_audit_metadata(target_headers, rows)
    identity_scope = f"{canonical.casefold()}#{sheet_name.casefold()}"
    final_headers, final_rows = ensure_row_metadata(target_headers, rows, kind, identity_scope=identity_scope)
    indexes = { _canonical_header(header): index for index, header in enumerate(final_headers) }
    return {
        (str(row[indexes["row id"]] or ""), str(row[indexes["content hash"]] or ""))
        for row in final_rows
        if row[indexes["row id"]] and row[indexes["content hash"]]
    }


def _verify_source_is_in_destination(source: Path, backup: Path, destination: Path, canonical: str) -> None:
    source_snapshot = _workbook_snapshot(source, canonical)
    backup_snapshot = _workbook_snapshot(backup, canonical)
    if source_snapshot != backup_snapshot:
        raise RuntimeError(f"The backup for {source.name} failed independent verification.")
    _validate_workbook(destination, canonical)
    if not _is_subset_workbook(source, destination):
        raise RuntimeError(f"The migrated records from {source.name} are missing in the destination.")
    source_ids = _expected_metadata_pairs(source, canonical)
    if source_ids and not source_ids.issubset(_metadata_pairs(destination)):
        raise RuntimeError(f"Row IDs or content hashes from {source.name} are missing in the destination.")
    _validate_destination_metadata(destination, canonical)


def _validate_destination_metadata(path: Path, canonical: str) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        found_data_sheet = False
        for sheet in workbook.worksheets:
            headers = [_canonical_header(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
            if "date" not in headers:
                continue
            found_data_sheet = True
            indexes = {header: index for index, header in enumerate(headers)}
            required = ("row id", "content hash", "created timestamp", "last updated timestamp", "updated device")
            if any(header not in indexes for header in required):
                raise RuntimeError(f"The destination metadata is incomplete for {Path(path).name}.")
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None for value in row):
                    continue
                if any(indexes[field] >= len(row) or not str(row[indexes[field]] or "").strip() for field in required):
                    raise RuntimeError(f"A migrated row in {Path(path).name} is missing its ID or history metadata.")
        if not found_data_sheet:
            raise RuntimeError(f"The destination data sheet is missing for {Path(path).name}.")
    finally:
        workbook.close()


def cleanup_migrated_root_sources(
    data_dir: Path,
    cleanup_root_dirs: list[Path],
    migration_report: dict,
    *,
    shared_sync_dir: Path | None = None,
) -> dict:
    """Remove only verified root-level legacy workbooks after migration commit.

    A durable ready-to-cleanup manifest is written before unlinking any source.
    Each deletion is recorded separately so launch recovery can safely resume
    if the process stops partway through cleanup.
    """
    data_dir = Path(data_dir).expanduser().resolve()
    roots = [Path(root).expanduser().resolve() for root in cleanup_root_dirs]
    if not roots:
        return {"status": "not-requested", "removed": []}
    if migration_report.get("complete") is not True:
        raise RuntimeError("Legacy root files were preserved because migration did not complete.")
    destinations = [data_dir / name for name in WORKBOOKS]
    if not all(path.is_file() and _is_valid(path, path.name) for path in destinations):
        raise RuntimeError("Legacy root files were preserved because all destination workbooks are not valid.")

    manifest_root = Path(shared_sync_dir or (data_dir.parent / "FinLedge Data Shared")).expanduser().resolve()
    manifest_path = manifest_root / ".finledge-legacy-cleanup.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if not isinstance(manifest, dict) or not isinstance(manifest.get("sources"), dict):
            manifest = {}
    except (OSError, ValueError):
        manifest = {}
    entries = manifest.setdefault("sources", {})
    reported_backups = migration_report.get("sourceBackups") or {}

    for root in roots:
        for canonical, aliases in WORKBOOKS.items():
            for filename in aliases:
                source = (root / filename).resolve()
                source_key = str(source)
                entry = entries.get(source_key)
                if source.is_file():
                    current_fingerprint = _source_identity_fingerprint(source)
                    backup_value = reported_backups.get(source_key)
                    is_new_source_version = entry is None or entry.get("fingerprint") != current_fingerprint or entry.get("state") in {"source_changed_preserved", "deleted"}
                    if is_new_source_version:
                        if not backup_value:
                            raise RuntimeError(f"Legacy root file {filename} has no verified migration backup; it was preserved.")
                        backup = Path(backup_value).expanduser().resolve()
                        _verify_source_is_in_destination(source, backup, data_dir / canonical, canonical)
                        entry = {
                            "canonical": canonical,
                            "fingerprint": current_fingerprint,
                            "backup": str(backup),
                            "state": "ready_to_cleanup",
                        }
                        entries[source_key] = entry
                        manifest.update({"version": 1, "complete": False, "updatedAt": datetime.now().astimezone().isoformat()})
                        _write_json_atomic(manifest_path, manifest)
                if not entry:
                    continue
                backup = Path(entry["backup"]).expanduser().resolve()
                destination = data_dir / str(entry["canonical"])
                shared_destination = manifest_root / str(entry["canonical"])
                if source.is_file():
                    if _source_identity_fingerprint(source) != entry.get("fingerprint"):
                        # A newer edit arrived after the backup. Keep it for a later migration pass.
                        entry["state"] = "source_changed_preserved"
                        continue
                    _verify_source_is_in_destination(source, backup, destination, str(entry["canonical"]))
                    # Root cleanup is committed only after the same records and
                    # metadata are independently present in the shared sync copy.
                    _verify_source_is_in_destination(source, backup, shared_destination, str(entry["canonical"]))
                    if _source_identity_fingerprint(source) != entry.get("fingerprint"):
                        entry["state"] = "source_changed_preserved"
                        manifest.update({"version": 1, "complete": False, "updatedAt": datetime.now().astimezone().isoformat()})
                        _write_json_atomic(manifest_path, manifest)
                        continue
                    source.unlink()
                elif entry.get("state") not in {"deleted", "source_changed_preserved"}:
                    # Covers a crash after unlink but before the manifest update.
                    _verify_source_is_in_destination(backup, backup, destination, str(entry["canonical"]))
                    _verify_source_is_in_destination(backup, backup, shared_destination, str(entry["canonical"]))
                entry["state"] = "deleted" if not source.exists() else "source_changed_preserved"
                entry["cleanedAt"] = datetime.now().astimezone().isoformat()
                manifest.update({"version": 1, "complete": False, "updatedAt": entry["cleanedAt"]})
                _write_json_atomic(manifest_path, manifest)

    requested_sources = [
        path for root in roots for aliases in WORKBOOKS.values() for name in aliases
        if (path := (root / name)).exists()
    ]
    unfinished = [entry for entry in entries.values() if entry.get("state") not in {"deleted", "source_changed_preserved"}]
    manifest.update({
        "version": 1,
        "complete": not unfinished and not requested_sources,
        "updatedAt": datetime.now().astimezone().isoformat(),
    })
    _write_json_atomic(manifest_path, manifest)
    removed = [path for path, entry in entries.items() if entry.get("state") == "deleted"]
    return {"status": "complete" if manifest["complete"] else "pending", "manifest": str(manifest_path), "removed": removed}


def _is_valid(path: Path, canonical: str | None = None) -> bool:
    try:
        candidate = canonical or Path(path).name
        _validate_workbook(path, candidate if candidate in REQUIRED_HEADER_GROUPS else None)
        return True
    except Exception:
        return False

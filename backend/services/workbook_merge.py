"""Shared row identity, tombstone, and deterministic workbook merge rules."""

from __future__ import annotations

import hashlib
import json
import os
import re
import threading
import time
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, uuid4, uuid5

from openpyxl import load_workbook
from .workbook_headers import header_key, map_workbook_row

ROW_ID_HEADER = "Row ID"
CONTENT_HASH_HEADER = "Content Hash"
DELETED_SHEET = "Deleted"
DELETED_HEADERS = ["Row ID", "Deleted Timestamp", "Deleted By Device"]
CONFLICTS_SHEET = "Conflicts"
CONFLICT_HEADERS = ["Row ID", "Losing Values", "Losing Device", "Timestamp"]
VALID_DEVICES = {"desktop-exe", "desktop-appx", "desktop-dev", "desktop-test", "mobile", "legacy"}
_identity_map_lock = threading.Lock()

BUSINESS_FIELDS = {
    "bank": ("Date", "Category", "Amount", "Description"),
    "share": ("Date", "Share Name", "Category", "Per Unit Price", "ASBA Charge", "Allotted", "Buy/Sell"),
    "personal_finance": ("Date", "Flow Type", "Direction", "Category", "Amount", "Description", "Source", "Source Ref"),
    "transfer": ("Date", "Flow Type", "Amount", "Description", "Source Ref"),
}


def get_device_label() -> str:
    """Return the current client label, with a safe development default."""
    value = str(os.environ.get("FINLEDGE_DEVICE_ID") or "desktop-dev").strip().lower()
    return value if value in VALID_DEVICES else "legacy"


def new_row_id() -> str:
    """Create the immutable UUID for a newly written transaction."""
    return str(uuid4())


def normalize_device_label(value: object, fallback: str = "legacy") -> str:
    normalized = str(value or "").strip().lower()
    if normalized == "desktop":
        return "legacy"
    return normalized if normalized in VALID_DEVICES else fallback


def kind_for_workbook(filename: str) -> str:
    name = Path(filename).name.casefold()
    if "share" in name:
        return "share"
    if "transfer" in name:
        return "transfer"
    if "personal_finance" in name:
        return "personal_finance"
    return "bank"


def canonical_header(value: object) -> str:
    return header_key(value)


def _normalize_value(header: str, value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    if header.casefold() in {"amount", "per unit price", "asba charge", "allotted"}:
        try:
            number = Decimal(text.replace(",", ""))
            if number.is_finite():
                return format(number.normalize(), "f")
        except (InvalidOperation, ValueError):
            pass
    if header.casefold() in {"category", "flow type", "direction", "buy/sell", "source"}:
        return re.sub(r"\s+", " ", text).casefold()
    return re.sub(r"\s+", " ", text)


def business_payload(headers: list[str], row: list[Any] | tuple[Any, ...], kind: str) -> dict[str, Any]:
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    payload = {}
    for field in BUSINESS_FIELDS[kind]:
        index = indexes.get(canonical_header(field))
        value = row[index] if index is not None and index < len(row) else None
        payload[field] = _normalize_value(field, value)
    return payload


def calculate_content_hash(headers: list[str], row: list[Any] | tuple[Any, ...], kind: str) -> str:
    encoded = json.dumps(business_payload(headers, row, kind), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha1(encoded.encode("utf-8")).hexdigest()


def ensure_identity_columns(headers: list[str]) -> list[str]:
    result = list(headers)
    for column in (ROW_ID_HEADER, CONTENT_HASH_HEADER):
        if column.casefold() not in {str(header).strip().casefold() for header in result}:
            result.append(column)
    return result


def map_row(source_headers: list[str], row: list[Any] | tuple[Any, ...], target_headers: list[str]) -> list[Any]:
    return map_workbook_row(source_headers, row, target_headers)


def _identity_map_path() -> Path | None:
    configured = str(os.environ.get("FINLEDGE_SHARED_SYNC_DIR") or "").strip()
    if not configured:
        return None
    return Path(configured).expanduser().resolve() / "row-identity-map.json"


def _read_identity_map(map_path: Path | None) -> dict[str, str]:
    if map_path is None:
        return {}
    try:
        mapping = json.loads(map_path.read_text(encoding="utf-8"))
        return mapping if isinstance(mapping, dict) else {}
    except (OSError, ValueError):
        return {}


def _persist_identity_map(map_path: Path | None, entries: dict[str, str]) -> None:
    if map_path is None:
        return

    map_path.parent.mkdir(parents=True, exist_ok=True)
    lock_path = map_path.with_suffix(".lock")
    deadline = time.monotonic() + 3
    with _identity_map_lock:
        descriptor = None
        while descriptor is None:
            try:
                descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            except FileExistsError:
                if time.monotonic() >= deadline:
                    return
                time.sleep(0.025)
        try:
            mapping = _read_identity_map(map_path)
            mapping.update({key: value for key, value in entries.items() if key not in mapping})
            staged = map_path.with_name(f".{map_path.name}.{os.getpid()}.tmp")
            try:
                staged.write_text(json.dumps(mapping, sort_keys=True, separators=(",", ":")), encoding="utf-8")
                os.replace(staged, map_path)
            finally:
                staged.unlink(missing_ok=True)
        finally:
            os.close(descriptor)
            lock_path.unlink(missing_ok=True)


def ensure_row_metadata(
    headers: list[str],
    rows: list[list[Any]],
    kind: str,
    *,
    identity_scope: str | None = None,
) -> tuple[list[str], list[list[Any]]]:
    """Add stable IDs and fresh business hashes, preserving existing valid IDs.

    The identity key uses a canonical workbook/sheet scope, business values,
    Created Timestamp, and an occurrence ordinal for exact duplicates. The
    ordinal identifies the duplicate set independent of row ordering; identical
    historical rows have no observable distinguishing attribute before IDs
    exist, so their assigned ID-to-physical-row pairing is intentionally not
    inferred.
    """
    headers = ensure_identity_columns(headers)
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    id_index = indexes["row id"]
    hash_index = indexes["content hash"]
    device_index = indexes.get("updated device")
    created_index = indexes.get("created timestamp")
    occurrences: Counter[str] = Counter()
    normalized_rows = []
    needs_identity_map = any(
        not str(row[id_index] if id_index < len(row) else "").strip()
        and normalize_device_label(row[device_index] if device_index is not None and device_index < len(row) else None) == "legacy"
        for row in rows
    )
    identity_map_path = _identity_map_path()
    identity_map = _read_identity_map(identity_map_path) if needs_identity_map else {}
    identity_entries: dict[str, str] = {}
    for original in rows:
        row = list(original) + [None] * max(0, len(headers) - len(original))
        digest = calculate_content_hash(headers, row, kind)
        created = row[created_index] if created_index is not None else ""
        identity_seed = json.dumps([identity_scope or kind, business_payload(headers, row, kind), _normalize_value("Created Timestamp", created)], ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        identity_digest = hashlib.sha1(identity_seed.encode("utf-8")).hexdigest()
        occurrence = occurrences[identity_digest]
        occurrences[identity_digest] += 1
        if not str(row[id_index] or "").strip():
            identity_key = f"{identity_digest}#{occurrence}"
            row[id_index] = str(identity_map.get(identity_key) or uuid5(NAMESPACE_URL, f"finledge-row-v1:{identity_key}"))
            identity_entries[identity_key] = str(row[id_index])
        row[hash_index] = digest
        if device_index is not None:
            row[device_index] = normalize_device_label(row[device_index])
        normalized_rows.append(row[:len(headers)])
    if identity_entries:
        _persist_identity_map(identity_map_path, identity_entries)
    return headers, normalized_rows


def populate_missing_audit_metadata(headers: list[str], rows: list[list[Any]]) -> list[list[Any]]:
    """Fill only absent historical audit fields without changing valid values."""
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    created_index = indexes.get("created timestamp")
    updated_index = indexes.get("last updated timestamp")
    device_index = indexes.get("updated device")
    date_index = indexes.get("date")
    populated = []
    for original in rows:
        row = list(original) + [None] * max(0, len(headers) - len(original))
        created = row[created_index] if created_index is not None else None
        if created_index is not None and not str(created or "").strip():
            historical_date = row[date_index] if date_index is not None else None
            if hasattr(historical_date, "isoformat"):
                created = historical_date.isoformat()
            else:
                created = str(historical_date or "").strip()
            row[created_index] = created or None
        if updated_index is not None and not str(row[updated_index] or "").strip():
            row[updated_index] = row[created_index] if created_index is not None else None
        if device_index is not None and not str(row[device_index] or "").strip():
            row[device_index] = "legacy"
        populated.append(row[:len(headers)])
    return populated


def refresh_sheet_metadata(sheet, kind: str, *, touched_row: int | None = None, device: str | None = None, identity_scope: str | None = None) -> None:
    headers = [str(sheet.cell(1, column).value or "") for column in range(1, sheet.max_column + 1)]
    canonical = {canonical_header(header): index for index, header in enumerate(headers)}
    id_index = canonical.get("row id")
    hash_index = canonical.get("content hash")
    if id_index is None or hash_index is None:
        raise ValueError("Workbook identity columns are missing; run the data migration first.")
    rows = [[sheet.cell(row, column + 1).value for column in range(len(headers))] for row in range(2, sheet.max_row + 1)]
    headers, rows = ensure_row_metadata(headers, rows, kind, identity_scope=identity_scope)
    device_index = {canonical_header(header): index for index, header in enumerate(headers)}.get("updated device")
    for offset, values in enumerate(rows, start=2):
        if touched_row == offset and device_index is not None:
            values[device_index] = device or get_device_label()
        for column, value in enumerate(values, start=1):
            sheet.cell(offset, column).value = value


def ensure_tombstone_and_conflict_sheets(workbook) -> None:
    for name, headers in ((DELETED_SHEET, DELETED_HEADERS), (CONFLICTS_SHEET, CONFLICT_HEADERS)):
        if name in workbook.sheetnames:
            sheet = workbook[name]
        else:
            sheet = workbook.create_sheet(name)
        for column, header in enumerate(headers, start=1):
            if not sheet.cell(1, column).value:
                sheet.cell(1, column).value = header


def append_tombstone(workbook, row_id: str, device: str | None = None, timestamp: str | None = None) -> None:
    ensure_tombstone_and_conflict_sheets(workbook)
    sheet = workbook[DELETED_SHEET]
    sheet.append([row_id, timestamp or datetime.now(timezone.utc).isoformat(timespec="seconds"), device or get_device_label()])


def read_tombstones(workbook) -> list[dict[str, str]]:
    if DELETED_SHEET not in workbook.sheetnames:
        return []
    sheet = workbook[DELETED_SHEET]
    headers = {canonical_header(cell.value): cell.column - 1 for cell in sheet[1]}
    result = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[headers.get("row id", 0)]:
            continue
        result.append({
            "row_id": str(row[headers.get("row id", 0)]),
            "deleted_timestamp": str(row[headers.get("deleted timestamp", 1)] or ""),
            "deleted_by_device": str(row[headers.get("deleted by device", 2)] or "legacy"),
        })
    return result


def read_conflicts(workbook) -> list[list[Any]]:
    if CONFLICTS_SHEET not in workbook.sheetnames:
        return []
    sheet = workbook[CONFLICTS_SHEET]
    return [[sheet.cell(row, column).value for column in range(1, len(CONFLICT_HEADERS) + 1)] for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value]


def _timestamp(row: list[Any], headers: list[str]) -> str:
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    index = indexes.get("last updated timestamp", indexes.get("created timestamp", -1))
    return str(row[index] or "") if 0 <= index < len(row) else ""


def _row_device(row: list[Any], headers: list[str]) -> str:
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    index = indexes.get("updated device")
    return normalize_device_label(row[index]) if index is not None and index < len(row) else "legacy"


def merge_row_sets(
    headers: list[str],
    current_rows: list[list[Any]],
    incoming_rows: list[list[Any]],
    kind: str,
    *,
    identity_scope: str | None = None,
    current_tombstones: list[dict[str, str]] | None = None,
    incoming_tombstones: list[dict[str, str]] | None = None,
    existing_conflicts: list[list[Any]] | None = None,
) -> dict[str, Any]:
    headers, current_rows = ensure_row_metadata(headers, current_rows, kind, identity_scope=identity_scope)
    _, incoming_rows = ensure_row_metadata(headers, incoming_rows, kind, identity_scope=identity_scope)
    indexes = {canonical_header(header): index for index, header in enumerate(headers)}
    id_index = indexes["row id"]
    hash_index = indexes["content hash"]

    def indexed(rows):
        return {str(row[id_index]): row for row in rows if str(row[id_index] or "").strip()}

    left = indexed(current_rows)
    right = indexed(incoming_rows)
    tombstones: dict[str, dict[str, str]] = {}
    for item in (current_tombstones or []) + (incoming_tombstones or []):
        row_id = str(item.get("row_id") or "").strip()
        if not row_id:
            continue
        previous = tombstones.get(row_id)
        if previous is None or str(item.get("deleted_timestamp") or "") > previous["deleted_timestamp"]:
            tombstones[row_id] = {
                "row_id": row_id,
                "deleted_timestamp": str(item.get("deleted_timestamp") or ""),
                "deleted_by_device": str(item.get("deleted_by_device") or "legacy"),
            }

    merged: dict[str, list[Any]] = {}
    conflicts = list(existing_conflicts or [])
    conflict_keys = {
        (str(row[0]), str(row[1]), str(row[2]))
        for row in conflicts if len(row) >= 3
    }
    report = {"added": 0, "updated": 0, "skipped_unchanged": 0, "conflicts": 0, "deleted": 0, "by_device": {}}
    for row_id in dict.fromkeys([*left.keys(), *right.keys()]):
        old, new = left.get(row_id), right.get(row_id)
        if old is not None and new is not None:
            if old[hash_index] == new[hash_index]:
                winner = old
                report["skipped_unchanged"] += 1
            else:
                old_ts, new_ts = _timestamp(old, headers), _timestamp(new, headers)
                old_rank = (old_ts, _row_device(old, headers), str(old[hash_index]))
                new_rank = (new_ts, _row_device(new, headers), str(new[hash_index]))
                winner, loser = (new, old) if new_rank > old_rank else (old, new)
                report["conflicts"] += 1
                report["updated"] += int(winner is new)
                report["skipped_unchanged"] += int(winner is old)
                values = {header: value for header, value in zip(headers, loser) if canonical_header(header) not in {"row id", "content hash", "created timestamp", "last updated timestamp", "updated device"}}
                losing_values = json.dumps(values, ensure_ascii=False, sort_keys=True)
                losing_device = _row_device(loser, headers)
                conflict_key = (row_id, losing_values, losing_device)
                if conflict_key not in conflict_keys:
                    conflicts.append([row_id, losing_values, losing_device, datetime.now(timezone.utc).isoformat(timespec="seconds")])
                    conflict_keys.add(conflict_key)
        elif new is not None:
            winner = new
            report["added"] += 1
        else:
            winner = old

        tombstone = tombstones.get(row_id)
        if tombstone and tombstone["deleted_timestamp"] > _timestamp(winner, headers):
            report["deleted"] += 1
            continue
        if tombstone:
            tombstones.pop(row_id, None)
        merged[row_id] = winner

    for row in merged.values():
        device = _row_device(row, headers)
        report["by_device"][device] = report["by_device"].get(device, 0) + 1
    return {
        "headers": headers,
        "rows": list(merged.values()),
        "tombstones": list(tombstones.values()),
        "conflicts": conflicts,
        "report": report,
        "message": format_merge_notice(report),
    }


def format_merge_notice(report: dict[str, Any]) -> str:
    if not any(report.get(key, 0) for key in ("added", "updated", "conflicts", "deleted")):
        return "Everything is up to date."
    parts = []
    if report.get("added"):
        parts.append(f"Added {report['added']}")
    if report.get("updated"):
        parts.append(f"updated {report['updated']}")
    if report.get("conflicts"):
        parts.append(f"kept the newer edit in {report['conflicts']} conflict{'s' if report['conflicts'] != 1 else ''}")
    if report.get("deleted"):
        parts.append(f"kept {report['deleted']} deleted")
    return ", ".join(parts) + "."


def recompute_derived(sheet, kind: str) -> None:
    headers = {canonical_header(cell.value): cell.column for cell in sheet[1]}
    if kind == "share":
        from .share_service import _recompute_sheet
        _recompute_sheet(sheet)
        return
    if kind == "bank":
        from .bank_service import _recompute_bank_sheet
        _recompute_bank_sheet(sheet)
    elif kind in {"personal_finance", "transfer"}:
        for row in range(2, sheet.max_row + 1):
            amount = float(sheet.cell(row, headers.get("amount", 5)).value or 0)
            direction = str(sheet.cell(row, headers.get("direction", 3)).value or "").strip().casefold()
            signed = amount if kind == "transfer" or direction == "income" else -abs(amount)
            sheet.cell(row, headers.get("signed amount", 6)).value = signed


def _find_data_sheet(workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    for sheet in workbook.worksheets:
        headers = {canonical_header(cell.value) for cell in sheet[1]}
        if "date" in headers and "amount" in headers:
            return sheet
    return workbook.active


def _sheet_rows(sheet) -> tuple[list[str], list[list[Any]]]:
    headers = [str(cell.value or "") for cell in sheet[1]]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2, values_only=False) if any(cell.value is not None for cell in row)]
    return headers, rows


def merge_workbook_files(source_path: Path, destination_path: Path, *, sheet_name: str, headers: list[str], kind: str, identity_scope: str | None = None) -> dict[str, Any]:
    """Merge source workbook into a caller-provided staging copy of destination."""
    source = load_workbook(source_path, data_only=False)
    destination = load_workbook(destination_path, data_only=False)
    try:
        source_sheet = _find_data_sheet(source, sheet_name)
        target_sheet = _find_data_sheet(destination, sheet_name)
        target_headers = ensure_identity_columns(headers)
        source_headers, raw_source_rows = _sheet_rows(source_sheet)
        target_source_headers, raw_target_rows = _sheet_rows(target_sheet)
        source_rows = [map_row(source_headers, row, target_headers) for row in raw_source_rows]
        target_rows = [map_row(target_source_headers, row, target_headers) for row in raw_target_rows]
        result = merge_row_sets(
            target_headers,
            target_rows,
            source_rows,
            kind,
            identity_scope=identity_scope or f"{kind}#{sheet_name.casefold()}",
            current_tombstones=read_tombstones(destination),
            incoming_tombstones=read_tombstones(source),
            existing_conflicts=read_conflicts(destination) + read_conflicts(source),
        )
        previous_tombstones = read_tombstones(destination)
        previous_conflicts = read_conflicts(destination)
        changed = any(result["report"].get(key, 0) for key in ("added", "updated", "conflicts", "deleted"))
        changed = changed or result["tombstones"] != previous_tombstones
        changed = changed or result["conflicts"] != previous_conflicts
        if target_sheet.title != sheet_name and sheet_name not in destination.sheetnames:
            target_sheet.title = sheet_name
        target_sheet.delete_rows(1, target_sheet.max_row)
        target_sheet.append(result["headers"])
        for row in result["rows"]:
            target_sheet.append(row)
        ensure_tombstone_and_conflict_sheets(destination)
        deleted_sheet = destination[DELETED_SHEET]
        deleted_sheet.delete_rows(2, max(0, deleted_sheet.max_row - 1))
        for tombstone in result["tombstones"]:
            deleted_sheet.append([tombstone["row_id"], tombstone["deleted_timestamp"], tombstone["deleted_by_device"]])
        conflict_sheet = destination[CONFLICTS_SHEET]
        conflict_sheet.delete_rows(2, max(0, conflict_sheet.max_row - 1))
        for conflict in result["conflicts"]:
            conflict_sheet.append(conflict)
        recompute_derived(target_sheet, kind)
        # Hashes exclude derived fields, so recomputing totals does not alter them.
        destination.save(destination_path)
        return {**result, "changed": bool(changed)}
    finally:
        source.close()
        destination.close()

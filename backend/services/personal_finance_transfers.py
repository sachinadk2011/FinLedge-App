"""Workbook storage and display shaping for Bank/Cash transfers."""

import uuid
from datetime import date
from decimal import Decimal
from typing import Optional

from openpyxl import load_workbook

from .excel_utils import commit_workbook_write, ensure_workbook_exists, iter_sheet_rows_as_dicts, workbook_for_update
from .personal_finance_summary import LINKED_MODULE_SOURCES, SHARE_ONLY_CATEGORIES, summarize_personal_finance_records
from .transfer_directions import (
    destination_flow,
    normalize_transfer_direction,
    origin_flow,
    transfer_flow_label,
)
from .workbook_merge import ensure_row_metadata, get_device_label, map_row, new_row_id, refresh_sheet_metadata
from .personal_finance_service import (
    DATA_DIR,
    HEADERS,
    TRANSFER_FILE_PATH,
    TRANSFER_SHEET_NAME,
    _current_timestamp,
    _file_lock,
    _to_float,
    read_personal_finance_records,
)


def _ensure_transfer_workbook_exists() -> None:
    def copy_legacy_transfers(workbook, sheet) -> bool:
        legacy_sheet = workbook["Personal Finance"] if "Personal Finance" in workbook.sheetnames else None
        if legacy_sheet is None:
            return False
        legacy_headers = [str(legacy_sheet.cell(row=1, column=column).value or "").strip() for column in range(1, legacy_sheet.max_column + 1)]
        legacy_rows = [
            map_row(legacy_headers, [legacy_sheet.cell(row=row, column=column).value for column in range(1, legacy_sheet.max_column + 1)], HEADERS)
            for row in range(2, legacy_sheet.max_row + 1)
        ]
        _, legacy_rows = ensure_row_metadata(HEADERS, legacy_rows, "transfer", identity_scope="personal_finance_transfer.xlsx#transfer flow")
        existing_rows = {
            tuple(str(sheet.cell(row=row, column=column).value or "") for column in range(1, len(HEADERS) + 1))
            for row in range(2, sheet.max_row + 1)
        }
        changed = False
        for values in legacy_rows:
            values[1] = normalize_transfer_direction(values[1]) or normalize_transfer_direction(values[2])
            if not values[1]:
                continue
            values[2] = "transfer"
            values[3] = "Transfer"
            values[7] = "transfer"
            signature = tuple(str(value or "") for value in values)
            if signature in existing_rows:
                continue
            sheet.append(values)
            existing_rows.add(signature)
            changed = True
        return changed

    ensure_workbook_exists(
        TRANSFER_FILE_PATH, TRANSFER_SHEET_NAME, HEADERS, DATA_DIR,
        legacy_sheet_copy=copy_legacy_transfers, rename_compatible_sheet=False,
    )

def create_transfer_record(
    entry_date: Optional[date],
    direction: str,
    amount: Decimal | float,
    description: Optional[str] = None,
) -> dict:
    """
    Create a single transfer record in personal_finance_transfer.xlsx.

    direction must be 'bank_to_cash' or 'cash_to_bank'.
    Validates that the transfer amount does not exceed the available income
    on the originating side before writing.
    """
    direction = normalize_transfer_direction(direction)
    if not direction:
        raise ValueError("Choose Bank to Cash or Cash to Bank.")

    amount_float = abs(float(amount))
    if amount_float <= 0:
        raise ValueError("Transfer amount must be greater than zero.")

    with workbook_for_update(TRANSFER_FILE_PATH, _file_lock, _ensure_transfer_workbook_exists) as workbook:
        _validate_available_income(direction, amount_float)

        entry_date = entry_date or date.today()
        timestamp = _current_timestamp()
        description_value = (description or "").strip()
        short_id = uuid.uuid4().hex[:8]
        transfer_ref = f"transfer:{entry_date.isoformat()}:{short_id}"

        sheet = workbook[TRANSFER_SHEET_NAME]
        sheet.append(
            [
                    entry_date.isoformat(),
                    direction,
                    "transfer",
                    "Transfer",
                    amount_float,
                    amount_float,
                    description_value,
                    "transfer",
                    timestamp,
                    timestamp,
                    transfer_ref,
                    get_device_label(),
                    new_row_id(),
                    None,
            ]
        )
        refresh_sheet_metadata(sheet, "transfer", touched_row=sheet.max_row, identity_scope="personal_finance_transfer.xlsx#transfer flow")

        return {
            "date": entry_date.isoformat(),
            "direction": direction,
            "category": "Transfer",
            "amount": amount_float,
            "description": description_value or None,
            "source": "transfer",
            "transfer_ref": transfer_ref,
            "timestamp": timestamp,
            "created_timestamp": timestamp,
            "last_updated_timestamp": timestamp,
            "updated_device": get_device_label(),
            "row_id": str(sheet.cell(sheet.max_row, 13).value or ""),
            "content_hash": str(sheet.cell(sheet.max_row, 14).value or ""),
            "file": str(TRANSFER_FILE_PATH),
        }

def read_transfer_records() -> list[dict]:
    """Read all raw transfer records from personal_finance_transfer.xlsx."""
    with _file_lock:
        _ensure_transfer_workbook_exists()
        return _read_transfer_records_unlocked()


def _read_transfer_records_unlocked() -> list[dict]:
    workbook = load_workbook(TRANSFER_FILE_PATH, data_only=True)
    try:
        sheet = workbook[TRANSFER_SHEET_NAME]
        records: list[dict] = []
        for row_idx, row in iter_sheet_rows_as_dicts(sheet):
            direction = normalize_transfer_direction(row.get("flow type"))
            if not direction:
                direction = normalize_transfer_direction(row.get("direction"))
            amount = abs(_to_float(row.get("amount")))
            created = row.get("created timestamp")
            updated = row.get("last updated timestamp") or created
            records.append(
                {
                    "id": f"transfer-{row_idx - 1}",
                    "display_id": f"T-{row_idx - 1}",
                    "date": str(row.get("date") or ""),
                    "transfer_direction": direction,
                    "direction": "transfer",
                    "category": "Transfer",
                    "amount": amount,
                    "description": str(row.get("description") or ""),
                    "source": "transfer",
                    "timestamp": str(updated or ""),
                    "created_timestamp": str(created or ""),
                    "last_updated_timestamp": str(updated or ""),
                    "source_ref": str(row.get("source ref") or ""),
                    "updated_device": str(row.get("updated device") or "legacy"),
                    "row_id": str(row.get("row id") or ""),
                    "content_hash": str(row.get("content hash") or ""),
                }
            )
        return records
    finally:
        workbook.close()


def _validate_available_income(direction: str, amount: float, excluding_transfer_id: int | None = None) -> float:
    flow = origin_flow(direction)
    records = read_personal_finance_records(flow_type=flow)
    if excluding_transfer_id is not None:
        excluded_id = f"transfer-{excluding_transfer_id}"
        records = [record for record in records if not (record.get("source") == "transfer" and record.get("id") == excluded_id)]
    available = summarize_personal_finance_records(records)[flow]["total_income"]
    if amount > available:
        side = flow.title()
        raise ValueError(f"Only NPR {max(0.0, available):,.2f} is available in {side}.")
    return available


def validate_transfer_import_rows(
    rows: list[list],
    headers: list[str],
    mode: str,
) -> None:
    """Replay transfer imports against dated income and existing transfers."""
    from .personal_finance_service import _read_personal_finance_records_for_flow

    with _file_lock:
        base_records = (
            _read_personal_finance_records_for_flow("bank")
            + _read_personal_finance_records_for_flow("cash")
        )
        transfers = read_transfer_records() if mode == "merge" else []
        header_indexes = {str(header or "").strip().casefold(): index for index, header in enumerate(headers)}

        events = []
        sequence = 0

        for record in base_records:
            source = str(record.get("source") or "manual").strip().casefold()
            category = str(record.get("category") or "").strip().casefold()
            if source in LINKED_MODULE_SOURCES or category in SHARE_ONLY_CATEGORIES:
                continue
            if str(record.get("direction") or "").strip().casefold() != "income":
                continue
            flow = str(record.get("flow_type") or "").strip().casefold()
            if flow not in {"bank", "cash"}:
                continue
            events.append((
                _transfer_event_order(
                    record.get("date"), record.get("created_timestamp") or record.get("timestamp"), 0, sequence
                ),
                "income",
                flow,
                abs(_to_float(record.get("amount"))),
            ))
            sequence += 1

        for transfer in transfers:
            direction = normalize_transfer_direction(transfer.get("transfer_direction"))
            if not direction:
                continue
            events.append((
                _transfer_event_order(
                    transfer.get("date"), transfer.get("created_timestamp") or transfer.get("timestamp"), 1, sequence
                ),
                direction,
                "",
                abs(_to_float(transfer.get("amount"))),
            ))
            sequence += 1

        for row in rows:
            def cell(name: str, default=None):
                index = header_indexes.get(name.casefold())
                return row[index] if index is not None and index < len(row) else default

            direction = normalize_transfer_direction(cell("Flow Type")) or normalize_transfer_direction(cell("Direction"))
            events.append((
                _transfer_event_order(
                    cell("Date"), cell("Created Timestamp") or cell("Last Updated Timestamp"), 1, sequence
                ),
                direction,
                "",
                abs(_to_float(cell("Amount", 0))),
            ))
            sequence += 1

        balances = {"bank": 0.0, "cash": 0.0}
        for _order, event_type, flow, amount in sorted(events, key=lambda event: event[0]):
            if event_type == "income":
                balances[flow] += amount
                continue

            direction = event_type
            origin = origin_flow(direction)
            destination = destination_flow(direction)
            available = round(balances[origin], 10)
            if amount > available:
                raise ValueError(
                    f"Transfer would overdraw {origin.title()}; only NPR {max(0.0, available):,.2f} is available."
                )
            balances[origin] = round(balances[origin] - amount, 10)
            balances[destination] = round(balances[destination] + amount, 10)


def _transfer_event_order(date_value, timestamp, priority: int, sequence: int) -> tuple[str, str, int, int]:
    date_key = str(date_value or "")[:10]
    timestamp_key = str(timestamp or "")
    return date_key, timestamp_key, priority, sequence


def update_transfer_record(
    transfer_id: int,
    entry_date: Optional[date],
    direction: str,
    amount: Decimal | float,
    description: Optional[str] = None,
) -> dict:
    """
    Update an existing transfer record (by sequential 1-based ID).
    Re-validates that the new amount does not exceed available origin income,
    excluding the old transfer's own contribution.
    """
    direction = normalize_transfer_direction(direction)
    if not direction:
        raise ValueError("Choose Bank to Cash or Cash to Bank.")

    amount_float = abs(float(amount))
    if amount_float <= 0:
        raise ValueError("Transfer amount must be greater than zero.")

    with _file_lock:
        _ensure_transfer_workbook_exists()
        workbook = load_workbook(TRANSFER_FILE_PATH)
        try:
            sheet = workbook[TRANSFER_SHEET_NAME]
            target_row = transfer_id + 1
            if target_row < 2 or target_row > sheet.max_row:
                raise ValueError(f"Transfer record {transfer_id} not found.")
            headers = {str(cell.value or "").strip().casefold(): cell.column for cell in sheet[1]}
            flow_type_column = headers.get("flow type", 2)
            direction_column = headers.get("direction", 3)
            old_direction = normalize_transfer_direction(
                sheet.cell(row=target_row, column=flow_type_column).value
            ) or normalize_transfer_direction(sheet.cell(row=target_row, column=direction_column).value)
            if not old_direction:
                raise ValueError(f"Transfer record {transfer_id} has an invalid direction.")
            workbook.close()

            # Excluding this row also removes its incoming effect when the new
            # origin is the side that previously received the transfer.
            _validate_available_income(direction, amount_float, excluding_transfer_id=transfer_id)

            workbook = load_workbook(TRANSFER_FILE_PATH)
            sheet = workbook[TRANSFER_SHEET_NAME]
            headers = {str(cell.value or "").strip().casefold(): cell.column for cell in sheet[1]}
            timestamp = _current_timestamp()
            entry_date_val = (entry_date or date.today()).isoformat()
            description_value = (description or "").strip()

            sheet.cell(row=target_row, column=headers.get("date", 1)).value = entry_date_val
            sheet.cell(row=target_row, column=headers.get("flow type", 2)).value = direction
            sheet.cell(row=target_row, column=headers.get("direction", 3)).value = "transfer"
            sheet.cell(row=target_row, column=headers.get("amount", 5)).value = amount_float
            sheet.cell(row=target_row, column=headers.get("signed amount", 6)).value = amount_float
            sheet.cell(row=target_row, column=headers.get("description", 7)).value = description_value
            sheet.cell(row=target_row, column=headers.get("last updated timestamp", 10)).value = timestamp
            sheet.cell(row=target_row, column=headers.get("updated device", 12)).value = get_device_label()
            refresh_sheet_metadata(sheet, "transfer", touched_row=target_row, identity_scope="personal_finance_transfer.xlsx#transfer flow")

            commit_workbook_write(workbook, TRANSFER_FILE_PATH)
        finally:
            workbook.close()
    return {
        "id": f"transfer-{transfer_id}",
        "date": entry_date_val,
        "direction": direction,
        "amount": amount_float,
        "description": description_value or None,
        "source": "transfer",
        "last_updated_timestamp": timestamp,
        "updated_device": get_device_label(),
        "row_id": str(sheet.cell(target_row, headers.get("row id", 13)).value or ""),
        "content_hash": str(sheet.cell(target_row, headers.get("content hash", 14)).value or ""),
    }

def _build_transfer_records_for_flow(flow_type: str) -> list[dict]:
    """
    Read transfer records and return them shaped for the given flow_type perspective.

    bank_to_cash transfer:
      - bank side: signed_amount = -amount  (bank gives money away → income reduced)
      - cash side: signed_amount = +amount  (cash receives money  → income added)

    cash_to_bank transfer:
      - cash side: signed_amount = -amount  (cash gives money away → income reduced)
      - bank side: signed_amount = +amount  (bank receives money  → income added)
    """
    raw_transfers = read_transfer_records()
    shaped: list[dict] = []
    for rec in raw_transfers:
        transfer_direction = normalize_transfer_direction(rec.get("transfer_direction"))
        amount = abs(_to_float(rec.get("amount")))
        if amount <= 0 or not transfer_direction:
            continue

        origin = origin_flow(transfer_direction)
        destination = destination_flow(transfer_direction)
        if flow_type not in {origin, destination}:
            continue
        signed = -amount if flow_type == origin else amount
        label = transfer_flow_label(transfer_direction, flow_type)

        desc = str(rec.get("description") or "").strip()
        shaped.append(
            {
                "id": rec["id"],
                "display_id": rec["display_id"],
                "date": rec["date"],
                "flow_type": flow_type,
                "direction": "transfer",
                "category": "Transfer",
                "amount": amount,
                "signed_amount": signed,
                "description": f"{label}{': ' + desc if desc else ''}",
                "source": "transfer",
                "transfer_direction": transfer_direction,
                "timestamp": rec.get("timestamp", ""),
                "created_timestamp": rec.get("created_timestamp", ""),
                "last_updated_timestamp": rec.get("last_updated_timestamp", ""),
                "source_ref": rec.get("source_ref", ""),
                "updated_device": rec.get("updated_device", "legacy"),
                "row_id": rec.get("row_id", ""),
                "content_hash": rec.get("content_hash", ""),
            }
        )
    return shaped


def _build_transfer_records_for_combined() -> list[dict]:
    """
    For the combined view: include each transfer exactly ONCE, shown from
    the originating side's perspective (negative signed_amount).
      bank_to_cash → shown under bank flow, label "Withdrawn to Cash"
      cash_to_bank → shown under cash flow, label "Deposited to Bank"
    """
    raw_transfers = read_transfer_records()
    shaped: list[dict] = []
    for rec in raw_transfers:
        transfer_direction = normalize_transfer_direction(rec.get("transfer_direction"))
        amount = abs(_to_float(rec.get("amount")))
        if amount <= 0 or not transfer_direction:
            continue

        flow_type = origin_flow(transfer_direction)
        label = transfer_flow_label(transfer_direction, flow_type)

        desc = str(rec.get("description") or "").strip()
        shaped.append(
            {
                "id": rec["id"],
                "display_id": rec["display_id"],
                "date": rec["date"],
                "flow_type": flow_type,
                "direction": "transfer",
                "category": "Transfer",
                "amount": amount,
                "signed_amount": -amount,  # originating side always loses
                "description": f"{label}{': ' + desc if desc else ''}",
                "source": "transfer",
                "transfer_direction": transfer_direction,
                "timestamp": rec.get("timestamp", ""),
                "created_timestamp": rec.get("created_timestamp", ""),
                "last_updated_timestamp": rec.get("last_updated_timestamp", ""),
                "source_ref": rec.get("source_ref", ""),
                "updated_device": rec.get("updated_device", "legacy"),
                "row_id": rec.get("row_id", ""),
                "content_hash": rec.get("content_hash", ""),
            }
        )
    return shaped

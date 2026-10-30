import os
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook

from backend.services import excel_utils


def _make_workbook(path, value="original"):
    workbook = Workbook()
    workbook.active.append(["Value"])
    workbook.active.append([value])
    workbook.save(path)
    workbook.close()


def test_snapshot_before_write_copies_live_workbook(tmp_path):
    live_file = tmp_path / "bank_transactions.xlsx"
    _make_workbook(live_file)

    snapshot = excel_utils.snapshot_before_write(live_file)

    assert snapshot is not None
    assert snapshot.parent == tmp_path / "backups" / "auto" / "bank_transactions"
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{12}", snapshot.stem)
    assert snapshot.exists()
    copied = load_workbook(snapshot, data_only=True)
    assert copied.active.cell(2, 1).value == "original"
    copied.close()


def test_snapshot_falls_back_to_active_data_parent_when_shared_backup_fails(tmp_path, monkeypatch):
    live_file = tmp_path / "Roaming" / "Finledge Testing AppX" / "FinLedge Data" / "share_transactions.xlsx"
    live_file.parent.mkdir(parents=True)
    _make_workbook(live_file)
    shared_root = tmp_path / "Documents" / "FinLedge Testing" / "backups" / "AppX"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared_root))
    real_copy = excel_utils.shutil.copy2

    def fail_shared_copy(source, destination, *args, **kwargs):
        if str(destination).startswith(str(shared_root)):
            raise FileNotFoundError(2, "simulated unavailable shared backup folder", str(destination))
        return real_copy(source, destination, *args, **kwargs)

    monkeypatch.setattr(excel_utils.shutil, "copy2", fail_shared_copy)

    snapshot = excel_utils.snapshot_before_write(live_file)

    expected_root = live_file.parent.parent / "backups"
    assert snapshot is not None
    assert snapshot.is_file()
    assert snapshot.is_relative_to(expected_root)
    saved = load_workbook(snapshot, data_only=True)
    assert saved.active.cell(2, 1).value == "original"
    saved.close()


def test_import_and_corrupted_backup_paths_use_short_timestamps_and_keep_filenames(tmp_path, monkeypatch):
    from backend.services.settings_data import workbook_import

    backup_root = tmp_path / "backups"
    monkeypatch.setattr(workbook_import, "get_backup_directory", lambda _path: backup_root)
    monkeypatch.setattr(excel_utils, "get_backup_directory", lambda _path: backup_root)
    live_file = tmp_path / "bank_transactions.xlsx"
    original_bytes = b"the exact prior file bytes"
    live_file.write_bytes(original_bytes)

    imported_backup = workbook_import._backup_live_file(live_file)

    assert imported_backup is not None
    assert imported_backup.name == live_file.name
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{12}", imported_backup.parent.name)
    assert imported_backup.read_bytes() == original_bytes

    workbook = excel_utils.repair_or_recover_workbook(live_file, "Bank", ["Value"])
    workbook.close()
    quarantine_files = list((backup_root / "corrupted").glob("*"))
    assert len(quarantine_files) == 1
    assert re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{12}", quarantine_files[0].name)
    assert (quarantine_files[0] / live_file.name).read_bytes() == original_bytes


def test_auto_backup_recovery_returns_user_warning(tmp_path):
    live_file = tmp_path / "cash_transactions.xlsx"
    backup_dir = tmp_path / "backups" / "auto" / live_file.stem
    backup_dir.mkdir(parents=True)
    backup_file = backup_dir / "20261022T000000.000000Z.xlsx"
    _make_workbook(backup_file, "recovered")
    live_file.write_bytes(b"not an xlsx workbook")

    recovered = excel_utils.safe_load_workbook(live_file, "Cash", ["Value"])
    assert recovered.active.cell(2, 1).value == "recovered"
    recovered.close()
    warnings = excel_utils.consume_recovery_warnings()
    assert any("recovered from a backup" in warning for warning in warnings)


def test_shared_backup_root_is_used_for_snapshot_and_restore(tmp_path, monkeypatch):
    live_file = tmp_path / "appx-working" / "bank_transactions.xlsx"
    live_file.parent.mkdir()
    _make_workbook(live_file, "shared-backup-value")
    shared_backup_root = tmp_path / "shared-sync" / "Backups"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared_backup_root))

    snapshot = excel_utils.snapshot_before_write(live_file)

    assert snapshot is not None
    assert snapshot.is_relative_to(shared_backup_root)
    live_file.write_bytes(b"damaged workbook")
    restored = excel_utils.safe_load_workbook(live_file, "Bank", ["Value"])
    try:
        assert restored.active.cell(2, 1).value == "shared-backup-value"
    finally:
        restored.close()
    assert excel_utils.consume_recovery_warnings()


def test_missing_workbook_starts_empty_without_restoring_stale_backup_or_warning(tmp_path):
    live_file = tmp_path / "bank_transactions.xlsx"
    backup_dir = tmp_path / "backups" / "auto" / live_file.stem
    backup_dir.mkdir(parents=True)
    backup_file = backup_dir / "20261022T000000.000000Z.xlsx"
    _make_workbook(backup_file, "stale data")

    workbook = excel_utils.safe_load_workbook(live_file, "Bank", ["Value"])
    assert workbook.active.max_row == 1
    workbook.close()
    assert backup_file.is_file()
    assert excel_utils.consume_recovery_warnings() == []

    saved = load_workbook(live_file, data_only=True)
    assert saved.active.max_row == 1
    saved.close()


def test_bank_add_update_delete_each_snapshot_prior_state(tmp_path, monkeypatch):
    from backend.services import bank_service

    live_file = tmp_path / "bank_transactions.xlsx"
    workbook = Workbook()
    workbook.active.title = bank_service.SHEET_NAME
    workbook.active.append(bank_service.HEADERS)
    workbook.create_sheet("Deleted").append(["Row ID", "Deleted Timestamp", "Deleted By Device"])
    workbook.create_sheet("Conflicts").append(["Row ID", "Losing Values", "Losing Device", "Timestamp"])
    workbook.save(live_file)
    workbook.close()
    monkeypatch.setattr(bank_service, "FILE_PATH", live_file)
    monkeypatch.setattr(bank_service, "DATA_DIR", tmp_path)
    backup_dir = tmp_path / "backups" / "auto" / live_file.stem

    bank_service.append_bank_record(None, "Interest Earned", 10.0, "Add test")
    first_snapshot = next(backup_dir.glob("*.xlsx"))
    first = load_workbook(first_snapshot, data_only=True)
    assert first.active.max_row == 1
    first.close()

    bank_service.update_bank_record(1, None, "Interest Earned", 12.0, "Update test")
    snapshots = sorted(backup_dir.glob("*.xlsx"), key=lambda path: path.stat().st_mtime)
    assert len(snapshots) == 2
    before_update = load_workbook(snapshots[1], data_only=True)
    assert before_update.active.cell(2, 3).value == 10.0
    before_update.close()

    bank_service.delete_bank_record(1)
    snapshots = sorted(backup_dir.glob("*.xlsx"), key=lambda path: path.stat().st_mtime)
    assert len(snapshots) == 3
    before_delete = load_workbook(snapshots[2], data_only=True)
    assert before_delete.active.cell(2, 3).value == 12.0
    before_delete.close()


def test_snapshot_retention_keeps_recent_daily_monthly_and_removes_old(tmp_path):
    snapshot_dir = tmp_path / "backups" / "auto" / "share_transactions"
    snapshot_dir.mkdir(parents=True)
    now = datetime.now(timezone.utc)

    def add_snapshot(name, age):
        path = snapshot_dir / name
        path.write_bytes(b"snapshot")
        timestamp = (now - age).timestamp()
        os.utime(path, (timestamp, timestamp))
        return path

    recent_one = add_snapshot("recent-one.xlsx", timedelta(hours=2))
    recent_two = add_snapshot("recent-two.xlsx", timedelta(hours=6))
    same_day = (now - timedelta(days=3)).replace(hour=12, minute=0, second=0, microsecond=0)
    daily_one = add_snapshot("daily-one.xlsx", now - same_day)
    daily_older_same_day = add_snapshot("daily-older.xlsx", now - (same_day - timedelta(hours=2)))
    monthly_one = add_snapshot("monthly-one.xlsx", timedelta(days=60))
    old = add_snapshot("old.xlsx", timedelta(days=400))

    excel_utils._prune_auto_snapshots(snapshot_dir, now)

    assert recent_one.exists() and recent_two.exists()
    assert daily_one.exists()
    assert not daily_older_same_day.exists()
    assert monthly_one.exists()
    assert not old.exists()


def test_snapshot_failure_aborts_bank_write(tmp_path, monkeypatch):
    from backend.services import bank_service

    live_file = tmp_path / "bank_transactions.xlsx"
    headers = bank_service.HEADERS
    workbook = Workbook()
    workbook.active.title = bank_service.SHEET_NAME
    workbook.active.append(headers)
    workbook.save(live_file)
    workbook.close()
    monkeypatch.setattr(bank_service, "FILE_PATH", live_file)
    monkeypatch.setattr(bank_service, "DATA_DIR", tmp_path)

    def fail_snapshot(_path):
        raise excel_utils.SnapshotError("Snapshot failed; write aborted.")

    monkeypatch.setattr(bank_service, "snapshot_before_write", fail_snapshot)

    try:
        bank_service.append_bank_record(None, "Interest Earned", 10.0)
    except excel_utils.SnapshotError as exc:
        assert "write aborted" in str(exc)
    else:
        raise AssertionError("Expected snapshot failure to abort the workbook write")

    saved = load_workbook(live_file, data_only=True)
    assert saved.active.max_row == 1
    saved.close()


def test_concurrent_bank_writes_keep_valid_snapshots(tmp_path, monkeypatch):
    from backend.services import bank_service

    live_file = tmp_path / "bank_transactions.xlsx"
    workbook = Workbook()
    workbook.active.title = bank_service.SHEET_NAME
    workbook.active.append(bank_service.HEADERS)
    workbook.create_sheet("Deleted").append(["Row ID", "Deleted Timestamp", "Deleted By Device"])
    workbook.create_sheet("Conflicts").append(["Row ID", "Losing Values", "Losing Device", "Timestamp"])
    workbook.save(live_file)
    workbook.close()
    monkeypatch.setattr(bank_service, "FILE_PATH", live_file)
    monkeypatch.setattr(bank_service, "DATA_DIR", tmp_path)

    with ThreadPoolExecutor(max_workers=6) as executor:
        futures = [
            executor.submit(bank_service.append_bank_record, None, "Interest Earned", i + 1.0, f"Concurrent {i}")
            for i in range(8)
        ]
        for future in futures:
            future.result()

    backup_dir = tmp_path / "backups" / "auto" / live_file.stem
    snapshots = list(backup_dir.glob("*.xlsx"))
    assert len(snapshots) == 8
    for snapshot in snapshots:
        saved = load_workbook(snapshot, data_only=True)
        assert saved.active.max_row >= 1
        saved.close()

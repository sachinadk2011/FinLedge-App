from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.services.workbook_merge import (
    calculate_content_hash,
    ensure_row_metadata,
    merge_row_sets,
    merge_workbook_files,
)


BANK_HEADERS = [
    "Date", "Category", "Amount", "Cumulative Amount", "Description",
    "Created Timestamp", "Last Updated Timestamp", "Updated Device", "Row ID", "Content Hash",
]


def bank_row(row_id: str, amount: float, updated: str, device: str) -> list[object]:
    row = ["2026-01-02", "Interest Earned", amount, 0, "interest", "2026-01-02T10:00:00", updated, device, row_id, None]
    row[-1] = calculate_content_hash(BANK_HEADERS, row, "bank")
    return row


def test_two_client_copies_merge_new_rows_and_unchanged_rows():
    left = bank_row("same", 10, "2026-01-02T10:00:00", "desktop-exe")
    right = bank_row("other", 15, "2026-01-03T10:00:00", "desktop-appx")

    result = merge_row_sets(BANK_HEADERS, [left], [left, right], "bank")

    assert {row[8] for row in result["rows"]} == {"same", "other"}
    assert result["report"]["added"] == 1
    assert result["report"]["skipped_unchanged"] == 1


def test_conflict_keeps_newest_and_archives_loser():
    older = bank_row("stable-id", 10, "2026-01-03T10:00:00", "desktop-exe")
    newer = bank_row("stable-id", 20, "2026-01-04T10:00:00", "desktop-appx")

    result = merge_row_sets(BANK_HEADERS, [older], [newer], "bank")

    assert result["rows"][0][2] == 20
    assert result["rows"][0][7] == "desktop-appx"
    assert result["report"]["conflicts"] == 1
    assert result["conflicts"][0][0] == "stable-id"
    assert result["conflicts"][0][2] == "desktop-exe"
    assert json.loads(result["conflicts"][0][1])["Amount"] == 10


def test_newer_tombstone_wins_and_newer_row_removes_tombstone():
    row = bank_row("deleted-id", 10, "2026-01-03T10:00:00", "desktop-exe")
    tombstone = {"row_id": "deleted-id", "deleted_timestamp": "2026-01-04T10:00:00", "deleted_by_device": "desktop-appx"}
    deleted = merge_row_sets(BANK_HEADERS, [row], [], "bank", current_tombstones=[tombstone])
    assert deleted["rows"] == []
    assert deleted["report"]["deleted"] == 1

    updated_row = bank_row("deleted-id", 11, "2026-01-05T10:00:00", "desktop-exe")
    restored = merge_row_sets(BANK_HEADERS, [updated_row], [], "bank", current_tombstones=[tombstone])
    assert len(restored["rows"]) == 1
    assert restored["tombstones"] == []


def test_business_hash_ignores_audit_timestamps():
    first = bank_row("hash-id", 10, "2026-01-03T10:00:00", "desktop-exe")
    second = list(first)
    second[5] = "2026-01-01T10:00:00"
    second[6] = "2026-02-01T10:00:00"
    second[7] = "desktop-appx"

    assert calculate_content_hash(BANK_HEADERS, first, "bank") == calculate_content_hash(BANK_HEADERS, second, "bank")


def test_legacy_rows_get_same_deterministic_ids_and_occurrence_ids():
    legacy_headers = BANK_HEADERS[:8]
    original = ["2026-01-02", "Interest Earned", 10, 0, "interest", "2026-01-02T10:00:00", "2026-01-02T10:00:00", "legacy"]
    rows = [original, list(original)]

    _, first = ensure_row_metadata(legacy_headers, rows, "bank")
    _, second = ensure_row_metadata(legacy_headers, rows, "bank")

    assert [row[8] for row in first] == [row[8] for row in second]
    assert first[0][8] != first[1][8]


def test_duplicate_legacy_records_get_unique_stable_ids_across_clients_and_reordering(tmp_path, monkeypatch):
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(tmp_path / "shared-sync"))
    legacy_headers = BANK_HEADERS[:8]
    duplicate = ["2026-01-02", "Interest Earned", 10, 0, "same", "2026-01-02T10:00:00", "2026-01-02T10:00:00", "legacy"]
    other = ["2026-01-03", "Interest Earned", 11, 11, "different", "2026-01-03T10:00:00", "2026-01-03T10:00:00", "legacy"]

    _, exe_rows = ensure_row_metadata(legacy_headers, [duplicate, list(duplicate), other], "bank", identity_scope="bank_transactions.xlsx#bank")
    _, appx_rows = ensure_row_metadata(legacy_headers, [other, list(duplicate), duplicate], "bank", identity_scope="bank_transactions.xlsx#bank")

    exe_by_business = {}
    appx_by_business = {}
    for target, rows in ((exe_by_business, exe_rows), (appx_by_business, appx_rows)):
        for row in rows:
            target.setdefault((row[0], row[1], row[2], row[4]), []).append(row[8])
    assert exe_by_business == appx_by_business
    assert len(set(exe_by_business[("2026-01-02", "Interest Earned", 10, "same")])) == 2
    assert all(exe_by_business[key] == appx_by_business[key] for key in exe_by_business)
    assert (tmp_path / "shared-sync" / "row-identity-map.json").is_file()


def _write_bank_workbook(path: Path, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bank"
    sheet.append(BANK_HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def test_file_merge_only_changes_staging_copy(tmp_path: Path):
    original_row = bank_row("original", 10, "2026-01-02T10:00:00", "desktop-exe")
    incoming_row = bank_row("incoming", 20, "2026-01-03T10:00:00", "desktop-appx")
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "destination.xlsx"
    stage = tmp_path / "stage.xlsx"
    _write_bank_workbook(source, [incoming_row])
    _write_bank_workbook(destination, [original_row])
    source_bytes = source.read_bytes()
    destination_bytes = destination.read_bytes()
    shutil.copy2(destination, stage)

    result = merge_workbook_files(source, stage, sheet_name="Bank", headers=BANK_HEADERS, kind="bank")

    assert result["report"]["added"] == 1
    assert source.read_bytes() == source_bytes
    assert destination.read_bytes() == destination_bytes
    merged = load_workbook(stage, read_only=True, data_only=True)
    try:
        assert merged["Bank"].max_row == 3
        assert {merged["Bank"].cell(row, 9).value for row in (2, 3)} == {"original", "incoming"}
        assert "Deleted" in merged.sheetnames and "Conflicts" in merged.sheetnames
    finally:
        merged.close()


def test_settings_merge_reimport_is_idempotent(tmp_path: Path, monkeypatch):
    from backend.services.settings_data.config import DATA_TYPE_CONFIG
    from backend.services.settings_data.workbook_import import import_data_file

    live_path = tmp_path / "live-bank.xlsx"
    monkeypatch.setitem(DATA_TYPE_CONFIG["bank"], "path", live_path)
    uploaded = tmp_path / "upload.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bank"
    sheet.append(["Date", "Category", "Amount", "Description"])
    sheet.append(["2026-01-02", "Interest Earned", 10, "same imported row"])
    workbook.save(uploaded)
    workbook.close()

    first = import_data_file("bank", uploaded, "replace")
    second = import_data_file("bank", uploaded, "merge")

    assert first["merge_report"]["added"] == 1
    assert second["merge_report"]["skipped_unchanged"] == 1
    merged = load_workbook(live_path, read_only=True, data_only=True)
    try:
        assert merged["Bank"].max_row == 2
        assert merged["Bank"].cell(2, 9).value
        assert merged["Bank"].cell(2, 10).value
        assert "Deleted" in merged.sheetnames and "Conflicts" in merged.sheetnames
    finally:
        merged.close()

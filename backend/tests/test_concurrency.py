"""
test_concurrency.py — Concurrent read/write stress test for FinLedge Excel storage.

Verifies:
1. Concurrent reads & writes against Bank, Share, and PF services do not corrupt
   .xlsx files or raise BadZipFile / InvalidFileException.
2. All appended records are saved without data loss.
3. Automatic backup recovery retrieves valid data if a file is damaged.
"""
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook

from backend.services.bank_service import append_bank_record, read_bank_records, FILE_PATH as BANK_FILE
from backend.services.share_service import append_share_record, read_share_records, FILE_PATH as SHARE_FILE
from backend.services.personal_finance_service import (
    append_personal_finance_record,
    read_personal_finance_records,
)
from backend.services.excel_utils import safe_load_workbook, find_most_recent_valid_backup


def test_concurrent_bank_reads_and_writes(isolated_data_dir):
    """Fire 20 concurrent thread operations at bank_service to test file locks."""
    num_writes = 10
    num_reads = 10

    def write_op(i):
        return append_bank_record(
            entry_date=date(2026, 1, 1),
            category="Interest Earned",
            amount=100.0 + i,
            description=f"Concurrent Test {i}",
        )

    def read_op(i):
        return read_bank_records()

    futures = []
    with ThreadPoolExecutor(max_workers=8) as executor:
        for i in range(num_writes):
            futures.append(executor.submit(write_op, i))
            futures.append(executor.submit(read_op, i))

        results = [f.result() for f in as_completed(futures)]

    assert len(results) == num_writes + num_reads

    # Confirm all 10 write records are present in the final workbook
    final_records = read_bank_records()
    descriptions = {r["description"] for r in final_records}
    for i in range(num_writes):
        assert f"Concurrent Test {i}" in descriptions

    # Validate openpyxl can load the file cleanly without BadZipFile
    wb = load_workbook(BANK_FILE, data_only=True)
    assert wb.active.max_row >= num_writes + 1
    wb.close()



def test_concurrent_multi_service_reads_and_writes(isolated_data_dir):
    """Simulate Summary page mount with concurrent requests across Bank, Share, and PF."""
    def bank_task():
        read_bank_records()
        return append_bank_record(date(2026, 2, 1), "Interest Tax", -15.0, "Tax test")

    def share_task():
        read_share_records()
        return append_share_record(date(2026, 2, 1), "AAPL", "buy", Decimal("150.0"), 5, "buy")

    def pf_task():
        return read_personal_finance_records("combined")

    futures = []
    with ThreadPoolExecutor(max_workers=10) as executor:
        for _ in range(5):
            futures.append(executor.submit(bank_task))
            futures.append(executor.submit(share_task))
            futures.append(executor.submit(pf_task))

        for f in as_completed(futures):
            _ = f.result()

    # Validate all workbooks load cleanly
    for file_path in (BANK_FILE, SHARE_FILE):
        wb = load_workbook(file_path, data_only=True)
        assert wb.active.max_row >= 1
        wb.close()


def test_automatic_backup_recovery(isolated_data_dir):
    """Confirm safe_load_workbook restores from valid backup when live file is corrupted."""
    # 1. Create a valid backup file in backups/import/
    backup_dir = isolated_data_dir / "backups" / "import" / "20260101-000000"
    backup_dir.mkdir(parents=True, exist_ok=True)
    valid_backup = backup_dir / "bank_transactions.xlsx"

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank"
    ws.append(["Date", "Category", "Amount", "Cumulative Amount", "Description", "Created Timestamp", "Last Updated Timestamp"])
    ws.append(["2026-01-01", "Interest Earned", 500.0, 500.0, "Backup Record 1", "2026-01-01T12:00:00", "2026-01-01T12:00:00"])
    wb.save(valid_backup)
    wb.close()

    # 2. Corrupt the live bank_transactions.xlsx file with bad text bytes
    with open(BANK_FILE, "wb") as f:
        f.write(b"CORRUPTED_ZIP_HEADER_BYTES_TEST")

    # 3. Call safe_load_workbook -> should quarantine corrupted file and recover from backup
    loaded_wb = safe_load_workbook(BANK_FILE, "Bank", ["Date", "Category", "Amount"])
    loaded_ws = loaded_wb.active

    # Check recovered row
    first_data_row = [loaded_ws.cell(2, col).value for col in range(1, 6)]
    loaded_wb.close()

    assert first_data_row[0] == "2026-01-01"
    assert first_data_row[1] == "Interest Earned"
    assert first_data_row[4] == "Backup Record 1"

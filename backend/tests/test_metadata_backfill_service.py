from __future__ import annotations

import hashlib
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from backend.services import bank_service
from backend.services.data_migration_service import migrate_sync_metadata_workbook_locked


def _write_bank(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bank"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _read_bank(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Bank"]
        headers = [str(sheet.cell(1, column).value or "") for column in range(1, sheet.max_column + 1)]
        rows = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]
        return headers, rows
    finally:
        workbook.close()


def _migrate(path: Path, monkeypatch):
    from backend.services.bank_service import HEADERS

    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(path.parent / "shared-sync"))
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(path.parent / "backups"))
    return migrate_sync_metadata_workbook_locked(path, "Bank", HEADERS, "bank")


def test_old_rows_receive_stable_ids_and_honest_legacy_metadata(tmp_path, monkeypatch):
    source = tmp_path / "exe" / "bank_transactions.xlsx"
    headers = ["Date", "Category", "Amount", "Description"]
    row = ["2020-04-05", "Interest Earned", 12.5, "old record"]
    _write_bank(source, headers, [row, list(row)])

    result = _migrate(source, monkeypatch)
    final_headers, final_rows = _read_bank(source)

    assert result["status"] == "migrated"
    assert len({entry[final_headers.index("Row ID")] for entry in final_rows}) == 2
    for entry in final_rows:
        assert entry[final_headers.index("Created Timestamp")] == "2020-04-05"
        assert entry[final_headers.index("Last Updated Timestamp")] == "2020-04-05"
        assert entry[final_headers.index("Updated Device")] == "legacy"
        assert entry[final_headers.index("Content Hash")]
    assert Path(result["backup_file"]).is_file()


def test_partial_upgrade_preserves_valid_id_timestamps_device_and_reuses_id(tmp_path, monkeypatch):
    source = tmp_path / "appx" / "bank_transactions.xlsx"
    headers = ["Date", "Category", "Amount", "Description", "Created Timestamp", "Last Updated Timestamp", "Updated Device", "Row ID"]
    row = ["2020-04-05", "Interest Earned", 12.5, "old record", "2019-01-02T03:04:05", "2024-03-04T05:06:07", "desktop-appx", "keep-this-id"]
    _write_bank(source, headers, [row])

    _migrate(source, monkeypatch)
    final_headers, rows = _read_bank(source)
    result = rows[0]

    assert result[final_headers.index("Row ID")] == "keep-this-id"
    assert result[final_headers.index("Created Timestamp")] == "2019-01-02T03:04:05"
    assert result[final_headers.index("Last Updated Timestamp")] == "2024-03-04T05:06:07"
    assert result[final_headers.index("Updated Device")] == "desktop-appx"
    assert result[final_headers.index("Content Hash")]


def test_independent_exe_and_appx_backfills_assign_same_ids_to_legacy_duplicates(tmp_path, monkeypatch):
    shared_sync = tmp_path / "shared-sync"
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared_sync))
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(tmp_path / "backups"))
    headers = ["Date", "Category", "Amount", "Description"]
    identical = ["2020-04-05", "Interest Earned", 12.5, "same"]
    exe = tmp_path / "exe" / "bank_transactions.xlsx"
    appx = tmp_path / "appx" / "bank_transactions.xlsx"
    _write_bank(exe, headers, [identical, list(identical), ["2021-02-03", "Interest Earned", 4, "other"]])
    _write_bank(appx, headers, [["2021-02-03", "Interest Earned", 4, "other"], list(identical), identical])

    _migrate(exe, monkeypatch)
    _migrate(appx, monkeypatch)
    exe_headers, exe_rows = _read_bank(exe)
    appx_headers, appx_rows = _read_bank(appx)
    id_index = exe_headers.index("Row ID")
    exe_ids = sorted(row[id_index] for row in exe_rows)
    appx_ids = sorted(row[appx_headers.index("Row ID")] for row in appx_rows)

    assert len(set(exe_ids)) == 3
    assert exe_ids == appx_ids


def test_failed_backup_keeps_original_workbook_unchanged(tmp_path, monkeypatch):
    import backend.services.data_migration_service as migration

    source = tmp_path / "bank_transactions.xlsx"
    _write_bank(source, ["Date", "Category", "Amount"], [["2020-04-05", "Interest Earned", 12.5]])
    original_bytes = source.read_bytes()
    monkeypatch.setattr(migration.shutil, "copy2", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("backup denied")))
    from backend.services.bank_service import HEADERS

    with pytest.raises(OSError, match="backup denied"):
        migrate_sync_metadata_workbook_locked(source, "Bank", HEADERS, "bank")
    assert hashlib.sha256(source.read_bytes()).hexdigest() == hashlib.sha256(original_bytes).hexdigest()


def test_locked_or_unreadable_workbook_is_preserved_then_retry_upgrades_it(tmp_path, monkeypatch):
    import backend.services.data_migration_service as migration

    source = tmp_path / "bank_transactions.xlsx"
    _write_bank(source, ["Date", "Category", "Amount"], [["2020-04-05", "Interest Earned", 12.5]])
    original_bytes = source.read_bytes()
    real_load = migration.load_workbook

    def deny_staged_reads(path, *args, **kwargs):
        if ".backfill-" in Path(path).name:
            raise PermissionError("simulated workbook lock")
        return real_load(path, *args, **kwargs)

    monkeypatch.setattr(migration, "load_workbook", deny_staged_reads)
    from backend.services.bank_service import HEADERS
    with pytest.raises(PermissionError, match="workbook lock"):
        migrate_sync_metadata_workbook_locked(source, "Bank", HEADERS, "bank")
    assert source.read_bytes() == original_bytes

    monkeypatch.setattr(migration, "load_workbook", real_load)
    assert migrate_sync_metadata_workbook_locked(source, "Bank", HEADERS, "bank")["status"] == "migrated"
    headers, rows = _read_bank(source)
    assert rows[0][headers.index("Row ID")]


def test_failed_atomic_commit_keeps_previous_workbook_and_can_resume(tmp_path, monkeypatch):
    import backend.services.data_migration_service as migration

    source = tmp_path / "bank_transactions.xlsx"
    _write_bank(source, ["Date", "Category", "Amount"], [["2020-04-05", "Interest Earned", 12.5]])
    original_bytes = source.read_bytes()
    real_replace = migration.os.replace

    def fail_commit(src, dst):
        if Path(dst) == source:
            raise PermissionError("simulated replacement failure")
        return real_replace(src, dst)

    monkeypatch.setattr(migration.os, "replace", fail_commit)
    from backend.services.bank_service import HEADERS
    with pytest.raises(PermissionError, match="replacement failure"):
        migrate_sync_metadata_workbook_locked(source, "Bank", HEADERS, "bank")
    assert source.read_bytes() == original_bytes

    monkeypatch.setattr(migration.os, "replace", real_replace)
    assert migrate_sync_metadata_workbook_locked(source, "Bank", HEADERS, "bank")["status"] == "migrated"
    assert _read_bank(source)[1][0][8]


def test_simultaneous_user_append_and_background_backfill_keep_both_rows(tmp_path, monkeypatch):
    source = tmp_path / "bank_transactions.xlsx"
    _write_bank(source, ["Date", "Category", "Amount", "Description"], [["2020-04-05", "Interest Earned", 12.5, "legacy"]])
    monkeypatch.setattr(bank_service, "DATA_DIR", tmp_path)
    monkeypatch.setattr(bank_service, "FILE_PATH", source)
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(tmp_path / "shared-sync"))
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(tmp_path / "backups"))
    from backend.services.bank_service import HEADERS

    def edit():
        return bank_service.append_bank_record(date(2026, 1, 1), "Interest Earned", 5, "concurrent entry")

    with ThreadPoolExecutor(max_workers=2) as pool:
        edit_future = pool.submit(edit)
        backfill_future = pool.submit(migrate_sync_metadata_workbook_locked, source, "Bank", HEADERS, "bank")
        edit_future.result(timeout=15)
        backfill_future.result(timeout=15)

    final_headers, rows = _read_bank(source)
    assert len(rows) == 2
    ids = [row[final_headers.index("Row ID")] for row in rows]
    assert all(ids)
    assert len(set(ids)) == 2
    assert any(row[final_headers.index("Description")] == "concurrent entry" for row in rows)

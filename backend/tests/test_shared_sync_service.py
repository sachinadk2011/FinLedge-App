from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.services.shared_sync_service import _shared_sync_lock, synchronize_data_directories
from backend.services.workbook_merge import calculate_content_hash


BANK_HEADERS = [
    "Date", "Category", "Amount", "Cumulative Amount", "Description",
    "Created Timestamp", "Last Updated Timestamp", "Updated Device", "Row ID", "Content Hash",
]


def _bank_row(row_id: str, amount: int, updated: str, device: str) -> list[object]:
    row = ["2026-01-01", "Interest Earned", amount, amount, row_id, "2026-01-01T00:00:00", updated, device, row_id, None]
    row[-1] = calculate_content_hash(BANK_HEADERS, row, "bank")
    return row


def _create_client_workbooks(directory: Path, bank_rows: list[list[object]]) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    layouts = {
        "bank_transactions.xlsx": ("Bank", BANK_HEADERS),
        "share_transactions.xlsx": ("Share", ["Date", "Share Name", "Category", "Per Unit Price", "ASBA Charge", "Allotted", "Buy/Sell", "Total Amount", "Profit/Loss", "Cumulative Profit", "Created Timestamp", "Last Updated Timestamp", "Updated Device", "Row ID", "Content Hash"]),
        "personal_finance_bank_flow.xlsx": ("Bank Flow", ["Date", "Flow Type", "Direction", "Category", "Amount", "Signed Amount", "Description", "Source", "Created Timestamp", "Last Updated Timestamp", "Source Ref", "Updated Device", "Row ID", "Content Hash"]),
        "personal_finance_cash_flow.xlsx": ("Cash Flow", ["Date", "Flow Type", "Direction", "Category", "Amount", "Signed Amount", "Description", "Source", "Created Timestamp", "Last Updated Timestamp", "Source Ref", "Updated Device", "Row ID", "Content Hash"]),
        "personal_finance_transfer.xlsx": ("Transfer Flow", ["Date", "Flow Type", "Direction", "Category", "Amount", "Signed Amount", "Description", "Source", "Created Timestamp", "Last Updated Timestamp", "Source Ref", "Updated Device", "Row ID", "Content Hash"]),
    }
    for filename, (sheet_name, headers) in layouts.items():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        if filename == "bank_transactions.xlsx":
            for row in bank_rows:
                sheet.append(row)
        workbook.save(directory / filename)
        workbook.close()


def _read_bank(path: Path) -> dict[str, tuple[int, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Bank"]
        return {
            str(sheet.cell(row, 9).value): (int(sheet.cell(row, 3).value), str(sheet.cell(row, 7).value))
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, 9).value
        }
    finally:
        workbook.close()


def test_sync_merges_both_client_copies_and_keeps_latest_conflict(tmp_path, monkeypatch):
    exe = tmp_path / "exe-working"
    appx = tmp_path / "appx-working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    shared.mkdir()
    _create_client_workbooks(exe, [_bank_row("exe-row", 10, "2026-01-02T00:00:00", "desktop-exe")])
    _create_client_workbooks(appx, [_bank_row("appx-row", 20, "2026-01-03T00:00:00", "desktop-appx")])
    _create_client_workbooks(shared, [_bank_row("exe-row", 11, "2026-01-04T00:00:00", "desktop-exe")])
    # Existing row ID with a newer AppX edit conflicts with the shared copy.
    appx_book = load_workbook(appx / "bank_transactions.xlsx")
    appx_sheet = appx_book["Bank"]
    appx_sheet.append(_bank_row("exe-row", 25, "2026-01-05T00:00:00", "desktop-appx"))
    appx_book.save(appx / "bank_transactions.xlsx")
    appx_book.close()

    report = synchronize_data_directories(exe, shared)
    synchronize_data_directories(appx, shared)
    synchronize_data_directories(exe, shared)

    assert report["status"] == "complete"
    for directory in (exe, shared):
        rows = _read_bank(directory / "bank_transactions.xlsx")
        assert rows["exe-row"][0] == 25
        assert rows["appx-row"][0] == 20
    workbook = load_workbook(shared / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert workbook["Conflicts"].max_row >= 2
    finally:
        workbook.close()
    journal = [json.loads(line) for line in (shared / "sync-journal.jsonl").read_text(encoding="utf-8").splitlines()]
    assert journal
    assert any(item["conflicts"] for item in journal)


def test_concurrent_clients_serialize_shared_sync_and_both_converge(tmp_path, monkeypatch):
    exe = tmp_path / "exe-working"
    appx = tmp_path / "appx-working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    _create_client_workbooks(exe, [_bank_row("exe-row", 10, "2026-01-02T00:00:00", "desktop-exe")])
    _create_client_workbooks(appx, [_bank_row("appx-row", 20, "2026-01-03T00:00:00", "desktop-appx")])
    _create_client_workbooks(shared, [])

    with _shared_sync_lock(shared) as acquired:
        assert acquired is True
        assert synchronize_data_directories(exe, shared)["status"] == "busy"
        assert not list(shared.glob(".sync-state-*.json"))
    for _ in range(2):
        synchronize_data_directories(exe, shared)
        synchronize_data_directories(appx, shared)

    expected_ids = {"exe-row", "appx-row"}
    assert set(_read_bank(exe / "bank_transactions.xlsx")) == expected_ids
    assert set(_read_bank(appx / "bank_transactions.xlsx")) == expected_ids
    assert set(_read_bank(shared / "bank_transactions.xlsx")) == expected_ids


def test_sync_is_idempotent_and_does_not_append_duplicate_journal_entries(tmp_path, monkeypatch):
    local = tmp_path / "working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    _create_client_workbooks(local, [_bank_row("stable", 10, "2026-01-02T00:00:00", "desktop-exe")])
    _create_client_workbooks(shared, [])

    synchronize_data_directories(local, shared)
    journal_path = shared / "sync-journal.jsonl"
    first_event_count = len(journal_path.read_text(encoding="utf-8").splitlines())
    first_rows = _read_bank(local / "bank_transactions.xlsx")
    synchronize_data_directories(local, shared)
    second_event_count = len(journal_path.read_text(encoding="utf-8").splitlines())

    assert first_rows == _read_bank(local / "bank_transactions.xlsx")
    assert first_event_count == second_event_count


def test_failed_sync_direction_does_not_advance_checkpoint_and_retries(tmp_path, monkeypatch):
    import backend.services.shared_sync_service as sync_service

    local = tmp_path / "working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    _create_client_workbooks(local, [_bank_row("retry-row", 10, "2026-01-02T00:00:00", "desktop-exe")])
    _create_client_workbooks(shared, [])
    state_path = shared / f".sync-state-{__import__('hashlib').sha256(str(local.resolve()).casefold().encode('utf-8')).hexdigest()[:16]}.json"

    original_merge_direction = sync_service._merge_direction
    for failed_status in ("retry", "source_missing"):
        def fail_second_direction(**kwargs):
            if kwargs["filename"] == "bank_transactions.xlsx" and kwargs["direction"] == "shared_to_local":
                return {"status": failed_status}
            return original_merge_direction(**kwargs)

        monkeypatch.setattr(sync_service, "_merge_direction", fail_second_direction)
        result = synchronize_data_directories(local, shared)
        assert result["status"] == "partial"
        assert result["files"][0]["status"] == "retry"
        state = json.loads(state_path.read_text(encoding="utf-8"))
        assert "bank_transactions.xlsx" not in state

    monkeypatch.setattr(sync_service, "_merge_direction", original_merge_direction)
    retried = synchronize_data_directories(local, shared)
    assert retried["status"] == "complete"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    assert "bank_transactions.xlsx" in state
    assert set(_read_bank(local / "bank_transactions.xlsx")) == {"retry-row"}


def test_missing_local_source_does_not_advance_its_checkpoint(tmp_path, monkeypatch):
    local = tmp_path / "working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    _create_client_workbooks(local, [])
    _create_client_workbooks(shared, [])
    missing_name = "share_transactions.xlsx"
    (local / missing_name).unlink()

    result = synchronize_data_directories(local, shared)
    assert result["status"] == "partial"
    missing_report = next(item for item in result["files"] if item["file"] == missing_name)
    assert missing_report["status"] == "local_missing"
    state_path = next(shared.glob(".sync-state-*.json"))
    assert missing_name not in json.loads(state_path.read_text(encoding="utf-8"))


def test_concurrent_edit_between_sync_directions_does_not_commit_checkpoint(tmp_path, monkeypatch):
    import backend.services.shared_sync_service as sync_service

    local = tmp_path / "working"
    shared = tmp_path / "shared"
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(shared / "Backups"))
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(shared))
    _create_client_workbooks(local, [_bank_row("original", 10, "2026-01-02T00:00:00", "desktop-exe")])
    _create_client_workbooks(shared, [])
    state_path = shared / f".sync-state-{__import__('hashlib').sha256(str(local.resolve()).casefold().encode('utf-8')).hexdigest()[:16]}.json"

    original_merge_direction = sync_service._merge_direction
    edited = {"done": False}

    def edit_between_directions(**kwargs):
        result = original_merge_direction(**kwargs)
        if kwargs["filename"] == "bank_transactions.xlsx" and kwargs["direction"] == "local_to_shared" and not edited["done"]:
            workbook = load_workbook(local / "bank_transactions.xlsx")
            workbook["Bank"].append(_bank_row("concurrent", 15, "2026-01-03T00:00:00", "desktop-exe"))
            workbook.save(local / "bank_transactions.xlsx")
            workbook.close()
            edited["done"] = True
        return result

    monkeypatch.setattr(sync_service, "_merge_direction", edit_between_directions)
    first = synchronize_data_directories(local, shared)
    assert first["status"] == "partial"
    assert "bank_transactions.xlsx" not in json.loads(state_path.read_text(encoding="utf-8"))

    monkeypatch.setattr(sync_service, "_merge_direction", original_merge_direction)
    retry = synchronize_data_directories(local, shared)
    assert retry["status"] == "complete"
    assert set(_read_bank(shared / "bank_transactions.xlsx")) == {"original", "concurrent"}

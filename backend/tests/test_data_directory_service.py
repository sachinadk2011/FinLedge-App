import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook

from backend.services.data_directory_service import (
    _copy_backup_history,
    backup_root_workbooks_before_data_folder,
    cleanup_migrated_root_sources,
    migrate_legacy_workbooks,
)
from backend.services import data_migration_service


def _workbook(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.append(["Date", "Category", "Amount"])
    workbook.active.append(["2026-01-01", value, 10])
    workbook.save(path)
    workbook.close()


def _create_remaining_valid_workbooks(destination):
    layouts = {
        "share_transactions.xlsx": ["Date", "Share Name", "Category"],
        "personal_finance_bank_flow.xlsx": ["Date", "Category", "Amount"],
        "personal_finance_cash_flow.xlsx": ["Date", "Category", "Amount"],
        "personal_finance_transfer.xlsx": ["Date", "Amount", "Direction"],
    }
    for filename, headers in layouts.items():
        workbook = Workbook()
        workbook.active.append(headers)
        workbook.save(destination / filename)
        workbook.close()


def _write_root_workbooks(root, marker="legacy"):
    layouts = {
        "bank_transactions.xlsx": (["Date", "Category", "Amount"], ["2026-01-01", f"{marker} bank", 10]),
        "share_transactions.xlsx": (["Date", "Share Name", "Category"], ["2026-01-01", "ABC", "buy"]),
        "personal_finance_bank_flow.xlsx": (["Date", "Category", "Amount"], ["2026-01-01", f"{marker} bank flow", 10]),
        "personal_finance_cash_flow.xlsx": (["Date", "Category", "Amount"], ["2026-01-01", f"{marker} cash flow", 10]),
        "personal_finance_transfer.xlsx": (["Date", "Amount", "Direction"], ["2026-01-01", 1, "bank_to_cash"]),
    }
    for filename, (headers, row) in layouts.items():
        workbook = Workbook()
        workbook.active.append(headers)
        workbook.active.append(row)
        workbook.save(root / filename)
        workbook.close()


def _commit_root_migration(root, data_dir, sync_dir):
    from backend.services.data_migration_service import run_pending_sync_metadata_backfill
    from backend.services.shared_sync_service import synchronize_data_directories

    previous = {name: os.environ.get(name) for name in ("FINLEDGE_SHARED_SYNC_DIR", "FINLEDGE_BACKUP_DIR")}
    os.environ["FINLEDGE_SHARED_SYNC_DIR"] = str(sync_dir)
    os.environ["FINLEDGE_BACKUP_DIR"] = str(sync_dir / "Backups")
    try:
        report = migrate_legacy_workbooks(
            data_dir,
            [root],
            sync_dir / "Backups",
            cleanup_root_dirs=[root],
            shared_sync_dir=sync_dir,
        )
        run_pending_sync_metadata_backfill(data_dir)
        synchronize_data_directories(data_dir, sync_dir)
        report = json.loads((data_dir / ".finledge-data-migration.json").read_text(encoding="utf-8"))
    finally:
        for name, value in previous.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
    return report, cleanup_migrated_root_sources(data_dir, [root], report, shared_sync_dir=sync_dir)


def test_root_workbooks_are_backed_up_byte_for_byte_before_missing_data_child_is_created(tmp_path):
    own_root = tmp_path / "Finledge EXE"
    data_child = own_root / "FinLedge Data"
    backup_root = tmp_path / "Documents" / "FinLedge" / "FinLedge Data" / "backups" / "EXE"
    names = (
        "bank_transactions.xlsx",
        "share_transactions.xlsx",
        "personal_finance_bank_flow.xlsx",
        "personal_finance_cash_flow.xlsx",
        "personal_finance_transfer.xlsx",
    )
    own_root.mkdir(parents=True)
    layouts = {
        "bank_transactions.xlsx": ["Date", "Category", "Amount"],
        "share_transactions.xlsx": ["Date", "Share Name", "Category"],
        "personal_finance_bank_flow.xlsx": ["Date", "Category", "Amount"],
        "personal_finance_cash_flow.xlsx": ["Date", "Category", "Amount"],
        "personal_finance_transfer.xlsx": ["Date", "Amount", "Direction"],
    }
    original_bytes = {}
    for name in names:
        workbook = Workbook()
        workbook.active.append(layouts[name])
        workbook.active.append(["2026-01-01", "synthetic legacy row", 10])
        workbook.save(own_root / name)
        workbook.close()
        original_bytes[name] = (own_root / name).read_bytes()

    backups = backup_root_workbooks_before_data_folder(data_child, backup_root)

    assert not data_child.exists()
    assert len(backups) == len(names)
    for backup in backups:
        backup_path = Path(backup)
        assert backup_path.read_bytes() == original_bytes[backup_path.name]
        validated = load_workbook(backup_path, read_only=True, data_only=True)
        try:
            assert validated.worksheets
            assert any(sheet.max_row == 2 for sheet in validated.worksheets)
        finally:
            validated.close()
    for name in names:
        assert (own_root / name).read_bytes() == original_bytes[name]


def test_migration_copies_valid_workbooks_and_leaves_sources(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "Documents" / "FinLedge"
    original = source / "bank_transactions.xlsx"
    _workbook(original, "keep")

    result = migrate_legacy_workbooks(destination, [source])

    assert (destination / original.name).is_file()
    assert original.is_file()
    assert result["status"] == "source_data_retained_for_review"
    assert json.loads((destination / ".finledge-data-migration.json").read_text())["complete"]


def test_migration_snapshot_falls_back_beside_active_data_when_primary_nested_backup_is_unavailable(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "Roaming" / "Finledge Testing AppX" / "FinLedge Data"
    backup_root = tmp_path / "Documents" / "FinLedge Testing" / "backups" / "AppX"
    (backup_root / "migration").parent.mkdir(parents=True)
    (backup_root / "migration").write_text("blocks nested snapshot creation")
    original = source / "bank_transactions.xlsx"
    _workbook(original, "keep after fallback")
    original_bytes = original.read_bytes()

    result = migrate_legacy_workbooks(destination, [source], backup_root)

    snapshot_dir = Path(result["migrationBackup"])
    assert snapshot_dir.is_dir()
    assert snapshot_dir.is_relative_to(destination.parent / "backups")
    assert (snapshot_dir / "src" / "1" / original.name).read_bytes() == original_bytes
    assert snapshot_dir.name.count(" ") == 1
    assert len(snapshot_dir.name) == 15  # YYYY-MM-DD HHMM
    assert (snapshot_dir / "src" / "1" / original.name).name == original.name
    assert (destination / original.name).is_file()
    assert original.read_bytes() == original_bytes


def test_backup_history_can_be_copied_to_fallback_backup_root(tmp_path, monkeypatch):
    from backend.services import data_directory_service

    legacy_root = tmp_path / "old"
    active_data = tmp_path / "Roaming" / "Finledge Testing AppX" / "FinLedge Data"
    configured_backup = tmp_path / "Documents" / "FinLedge Testing" / "backups" / "AppX"
    fallback_backup = active_data.parent / "backups"
    legacy_backup = legacy_root / "backups" / "sync-schema-v1" / "old tag" / "bank_transactions.xlsx"
    legacy_backup.parent.mkdir(parents=True)
    legacy_backup.write_bytes(b"backup history stays intact")
    copy_verified = data_directory_service._copy_verified

    def deny_configured_backup_root(source, destination):
        if Path(destination).is_relative_to(configured_backup):
            raise OSError("configured AppX backup path is not writable")
        return copy_verified(source, destination)

    monkeypatch.setattr(data_directory_service, "_copy_verified", deny_configured_backup_root)

    copied = _copy_backup_history(
        active_data,
        [legacy_root],
        configured_backup,
    )

    copied_path = Path(copied[0])
    assert copied_path.is_relative_to(fallback_backup)
    assert copied_path.read_bytes() == legacy_backup.read_bytes()
    assert legacy_backup.read_bytes() == b"backup history stays intact"


def test_migration_reads_workbooks_from_nested_testing_exe_data_folder(tmp_path):
    legacy_root = tmp_path / "AppData" / "Roaming" / "Finledge Testing EXE"
    source = legacy_root / "FinLedge Data"
    destination = tmp_path / "shared" / "FinLedge Data"
    original = source / "bank_transactions.xlsx"
    _workbook(original, "testing exe legacy row")

    result = migrate_legacy_workbooks(destination, [legacy_root])

    assert (destination / original.name).is_file()
    assert result["files"][0]["source"] == str(original.resolve())
    assert original.is_file()


def test_partial_legacy_workbooks_outside_data_folder_migrate_from_old_root_and_initialize_missing_files(tmp_path, monkeypatch):
    from backend.services import bank_service, personal_finance_service, personal_finance_transfers, share_service
    from backend.services.workbook_initialization import WORKBOOK_NAMES, initialize_data_workbooks

    old_root = tmp_path / "AppData" / "Roaming" / "financial-tracker-desktop"
    destination = tmp_path / "AppData" / "Roaming" / "Finledge Testing EXE" / "FinLedge Data"
    old_bank_file = old_root / "bank.xlsx"
    _workbook(old_bank_file, "old root bank row")
    original_bytes = old_bank_file.read_bytes()

    result = migrate_legacy_workbooks(destination, [old_root])

    migrated = destination / "bank_transactions.xlsx"
    assert migrated.is_file()
    assert result["files"][0]["status"] == "copied_from_legacy"
    assert result["files"][0]["source"] == str(old_bank_file.resolve())
    assert old_bank_file.read_bytes() == original_bytes
    assert sum((destination / name).is_file() for name in WORKBOOK_NAMES) == 1

    monkeypatch.setattr(bank_service, "DATA_DIR", destination)
    monkeypatch.setattr(bank_service, "FILE_PATH", migrated)
    monkeypatch.setattr(share_service, "DATA_DIR", destination)
    monkeypatch.setattr(share_service, "FILE_PATH", destination / "share_transactions.xlsx")
    monkeypatch.setattr(personal_finance_service, "DATA_DIR", destination)
    monkeypatch.setattr(personal_finance_service, "BANK_FLOW_FILE_PATH", destination / "personal_finance_bank_flow.xlsx")
    monkeypatch.setattr(personal_finance_service, "CASH_FLOW_FILE_PATH", destination / "personal_finance_cash_flow.xlsx")
    monkeypatch.setattr(personal_finance_service, "TRANSFER_FILE_PATH", destination / "personal_finance_transfer.xlsx")
    monkeypatch.setattr(personal_finance_transfers, "DATA_DIR", destination)
    monkeypatch.setattr(personal_finance_transfers, "TRANSFER_FILE_PATH", destination / "personal_finance_transfer.xlsx")

    initialized = initialize_data_workbooks(destination)

    assert initialized["status"] == "initialized"
    assert all((destination / name).is_file() for name in WORKBOOK_NAMES)
    assert old_bank_file.read_bytes() == original_bytes
    assert any(record["category"] == "old root bank row" for record in bank_service.read_bank_records())
    bank = load_workbook(migrated, read_only=True, data_only=True)
    try:
        assert bank["Bank"].cell(2, 2).value == "old root bank row"
    finally:
        bank.close()


def test_migration_rerun_is_noop(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    _workbook(source / "bank_transactions.xlsx", "first")
    migrate_legacy_workbooks(destination, [source])
    _create_remaining_valid_workbooks(destination)
    original_destination = (destination / "bank_transactions.xlsx").read_bytes()

    result = migrate_legacy_workbooks(destination, [source])

    assert result["alreadyCompleted"] is True
    assert (destination / "bank_transactions.xlsx").read_bytes() == original_destination
    assert (source / "bank_transactions.xlsx").is_file()


def test_new_legacy_source_changes_fingerprint_and_is_merged(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    _workbook(source / "bank_transactions.xlsx", "first")
    first = migrate_legacy_workbooks(destination, [source])
    _create_remaining_valid_workbooks(destination)
    _workbook(source / "share_transactions.xlsx", "new share")

    second = migrate_legacy_workbooks(destination, [source])

    assert first["sourceFingerprint"] != second["sourceFingerprint"]
    assert (destination / "share_transactions.xlsx").is_file()


def test_existing_workbook_rows_are_merged_without_overwriting(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    _workbook(source / "bank_transactions.xlsx", "source")
    _workbook(destination / "bank_transactions.xlsx", "existing")

    result = migrate_legacy_workbooks(destination, [source])

    assert result["status"] == "source_data_retained_for_review"
    workbook = load_workbook(destination / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    assert [row[:3] for row in rows] == [("2026-01-01", "existing", 10), ("2026-01-01", "source", 10)]
    assert all(row[8] and row[9] for row in rows)
    assert (source / "bank_transactions.xlsx").is_file()


def test_empty_destination_is_backed_up_then_replaced_with_legacy_data(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    empty = destination / "bank_transactions.xlsx"
    _workbook(empty, "")
    # Keep a header-only, valid workbook to model an earlier app launch.
    workbook = load_workbook(empty)
    workbook.active.delete_rows(2, 1)
    workbook.save(empty)
    workbook.close()
    legacy = source / "bank_transactions.xlsx"
    _workbook(legacy, "legacy row")
    legacy_bytes = legacy.read_bytes()

    result = migrate_legacy_workbooks(destination, [source])

    copied = load_workbook(destination / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert list(copied["Bank"].iter_rows(min_row=2, values_only=True)) == [("2026-01-01", "legacy row", 10)]
    finally:
        copied.close()
    assert result["files"][0]["status"] == "replaced_empty_destination"
    assert (Path(result["migrationBackup"]) / "dest" / "bank_transactions.xlsx").is_file()
    assert legacy.read_bytes() == legacy_bytes


def test_old_headers_merge_by_name_and_fill_audit_columns(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    old_path = source / "bank_transactions.xlsx"
    old_path.parent.mkdir(parents=True)
    old_book = Workbook()
    old_sheet = old_book.active
    old_sheet.title = "Old Bank"
    old_sheet.append(["Date", "Category", "Amount", "Timestamp"])
    old_sheet.append(["2026-02-01", "old entry", 14, "2026-02-01 09:00"])
    old_book.save(old_path)
    old_book.close()
    old_bytes = old_path.read_bytes()

    new_book = Workbook()
    new_sheet = new_book.active
    new_sheet.title = "Bank"
    new_sheet.append(["Date", "Category", "Amount", "Created Timestamp", "Last Updated Timestamp", "Updated Device"])
    new_sheet.append(["2026-01-01", "existing entry", 9, "created", "updated", "desktop"])
    destination.mkdir(parents=True)
    new_book.save(destination / "bank_transactions.xlsx")
    new_book.close()

    migrate_legacy_workbooks(destination, [source])

    merged = load_workbook(destination / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert "Bank" in merged.sheetnames
        rows = list(merged["Bank"].iter_rows(min_row=2, values_only=True))
    finally:
        merged.close()
    assert [row[:8] for row in rows] == [
        ("2026-01-01", "existing entry", 9, 9, None, "created", "updated", "legacy"),
        ("2026-02-01", "old entry", 14, 23, None, "2026-02-01 09:00", "2026-02-01 09:00", "legacy"),
    ]
    assert all(row[8] and row[9] for row in rows)
    assert old_path.read_bytes() == old_bytes


def test_backup_is_not_restored_for_a_workbook_that_was_absent(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    _workbook(source / "backups" / "auto" / "bank_transactions" / "snapshot.xlsx", "stale backup")

    result = migrate_legacy_workbooks(destination, [source])

    assert not (destination / "bank_transactions.xlsx").exists()
    assert result["files"][0]["status"] == "not_present"


def test_migration_marks_a_fresh_data_folder_as_new_install(tmp_path):
    destination = tmp_path / "new"

    result = migrate_legacy_workbooks(destination, [tmp_path / "no-old-data"])

    assert result["status"] == "new_install"
    assert all(item["status"] == "not_present" for item in result["files"])

    _workbook(destination / "bank_transactions.xlsx", "first entry")
    _create_remaining_valid_workbooks(destination)
    next_start = migrate_legacy_workbooks(destination, [])
    assert next_start["status"] == "existing_data"
    assert next_start["alreadyCompleted"] is True


def test_migration_preserves_and_validates_backup_history(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "new"
    backup = source / "backups" / "auto" / "bank_transactions" / "snapshot.xlsx"
    _workbook(backup, "backup")

    result = migrate_legacy_workbooks(destination, [source])

    assert (destination / "backups" / "auto" / "bank_transactions" / "snapshot.xlsx").is_file()
    assert backup.is_file()
    assert result["copiedBackupFiles"]


def test_valid_legacy_source_is_never_removed_and_is_merged_once(tmp_path):
    source = tmp_path / "old"
    destination = tmp_path / "Documents" / "FinLedge"
    original = source / "bank_transactions.xlsx"
    original.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Category", "Amount"])
    sheet.append(["2026-01-01", "Interest Earned", 25])
    workbook.save(original)
    workbook.close()

    first = migrate_legacy_workbooks(destination, [source])
    original_bytes = original.read_bytes()
    second = migrate_legacy_workbooks(destination, [source])

    copied = load_workbook(destination / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert list(copied.active.iter_rows(min_row=2, values_only=True)) == [("2026-01-01", "Interest Earned", 25)]
    finally:
        copied.close()
    assert original.is_file()
    assert original.read_bytes() == original_bytes
    assert str(original) in first["retainedLegacyFiles"]
    assert second["migrationBackup"] is None


def test_schema_migration_marker_skips_workbook_loads_when_unchanged(tmp_path, monkeypatch):
    data_dir = tmp_path / "isolated-data"
    first = data_migration_service.run_pending_data_migrations(data_dir)
    assert first["status"] == "not-needed"
    assert not first.get("skipped")

    def unexpected_workbook_load(*_args, **_kwargs):
        raise AssertionError("unchanged startup must not open a workbook")

    monkeypatch.setattr(data_migration_service, "load_workbook", unexpected_workbook_load)
    second = data_migration_service.run_pending_data_migrations(data_dir)

    assert second["status"] == "not-needed"
    assert second["skipped"] is True


def test_root_cleanup_waits_for_verified_backup_and_destination_and_never_removes_unrelated_files(tmp_path):
    root = tmp_path / "Finledge Testing EXE"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    _write_root_workbooks(root)
    unrelated = root / "notes.txt"
    unrelated.write_text("keep", encoding="utf-8")
    original_hashes = {path.name: __import__("hashlib").sha256(path.read_bytes()).hexdigest() for path in root.glob("*.xlsx")}

    report, cleanup = _commit_root_migration(root, data_dir, sync_dir)

    assert report["complete"] is True
    assert cleanup["status"] == "complete"
    assert not list(root.glob("*.xlsx"))
    assert unrelated.read_text(encoding="utf-8") == "keep"
    for filename, backup_name in report["sourceBackups"].items():
        source = Path(filename)
        backup = Path(backup_name)
        assert backup.is_file()
        assert __import__("hashlib").sha256(backup.read_bytes()).hexdigest() == original_hashes[source.name]
        wb = load_workbook(backup, read_only=True, data_only=True)
        try:
            assert wb.worksheets
        finally:
            wb.close()
    bank = load_workbook(data_dir / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert bank["Bank"].max_row == 2
        assert bank["Bank"].cell(2, 9).value
        assert bank["Bank"].cell(2, 10).value
    finally:
        bank.close()

    second = migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    assert all((data_dir / item["name"]).is_file() for item in second["files"])
    bank = load_workbook(data_dir / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert bank["Bank"].max_row == 2
    finally:
        bank.close()


def test_root_cleanup_preserves_sources_until_shared_copy_is_validated(tmp_path, monkeypatch):
    import pytest
    from zipfile import BadZipFile
    from backend.services.data_migration_service import run_pending_sync_metadata_backfill
    from backend.services.shared_sync_service import synchronize_data_directories

    for shared_state in ("missing", "corrupt"):
        root = tmp_path / shared_state / "Finledge Testing EXE"
        data_dir = root / "FinLedge Data"
        sync_dir = root / "FinLedge Data Shared"
        root.mkdir(parents=True)
        _write_root_workbooks(root)
        monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(sync_dir))
        monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(sync_dir / "Backups"))

        report = migrate_legacy_workbooks(
            data_dir,
            [root],
            sync_dir / "Backups",
            cleanup_root_dirs=[root],
            shared_sync_dir=sync_dir,
        )
        run_pending_sync_metadata_backfill(data_dir)
        if shared_state == "corrupt":
            synchronize_data_directories(data_dir, sync_dir)
            (sync_dir / "bank_transactions.xlsx").write_bytes(b"corrupted shared copy")

        expected_error = FileNotFoundError if shared_state == "missing" else BadZipFile
        with pytest.raises(expected_error):
            cleanup_migrated_root_sources(data_dir, [root], report, shared_sync_dir=sync_dir)
        assert len(list(root.glob("*.xlsx"))) == len(report["sourceBackups"])


def test_backup_creation_failure_keeps_root_files_and_does_not_start_migration(tmp_path, monkeypatch):
    from backend.services import data_directory_service

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    _workbook(root / "bank_transactions.xlsx", "legacy")
    original = data_directory_service._copy_verified_workbook

    def fail_backup(source, destination, canonical):
        if Path(source).parent == root:
            raise IOError("simulated backup failure")
        return original(source, destination, canonical)

    monkeypatch.setattr(data_directory_service, "_copy_verified_workbook", fail_backup)
    import pytest
    with pytest.raises(RuntimeError, match="Could not safely prepare"):
        migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)

    assert (root / "bank_transactions.xlsx").is_file()
    assert not (data_dir / "bank_transactions.xlsx").exists()


def test_corrupt_backup_is_rejected_before_destination_copy_and_preserves_source(tmp_path, monkeypatch):
    from backend.services import data_directory_service
    import pytest

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    source = root / "bank_transactions.xlsx"
    _workbook(source, "legacy")
    real_copy = data_directory_service._copy_verified

    def corrupt_staged_backup(original, destination):
        real_copy(original, destination)
        if "\\src\\" in str(destination).casefold():
            Path(destination).write_bytes(b"corrupted backup")

    monkeypatch.setattr(data_directory_service, "_copy_verified", corrupt_staged_backup)
    with pytest.raises(RuntimeError, match="Could not safely prepare"):
        migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)

    assert source.is_file()
    assert not (data_dir / source.name).exists()


def test_destination_copy_failure_keeps_root_source(tmp_path, monkeypatch):
    from backend.services import data_directory_service
    import pytest

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    source = root / "bank_transactions.xlsx"
    _workbook(source, "legacy")
    real_copy = data_directory_service._copy_verified

    def fail_destination_copy(original, destination):
        if Path(destination).parent == data_dir and Path(destination).name.startswith(".migration-"):
            raise IOError("simulated destination copy failure")
        return real_copy(original, destination)

    monkeypatch.setattr(data_directory_service, "_copy_verified", fail_destination_copy)
    with pytest.raises(RuntimeError, match="Could not safely prepare"):
        migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)

    assert source.is_file()
    assert not (data_dir / source.name).exists()


def test_destination_validation_failure_preserves_root_source(tmp_path, monkeypatch):
    from backend.services import data_directory_service
    import pytest

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    source = root / "bank_transactions.xlsx"
    _workbook(source, "legacy")
    real_valid = data_directory_service._is_valid

    def reject_destination(path, canonical=None):
        if Path(path).parent == data_dir and Path(path).name == source.name:
            return False
        return real_valid(path, canonical)

    monkeypatch.setattr(data_directory_service, "_is_valid", reject_destination)
    with pytest.raises(RuntimeError, match="Could not safely prepare"):
        migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)

    assert source.is_file()


def test_restart_after_crash_between_verified_backup_and_migration_completes_safely(tmp_path, monkeypatch):
    from backend.services import data_directory_service
    import pytest

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    source = root / "bank_transactions.xlsx"
    _workbook(source, "legacy")
    original_validate = data_directory_service._validate_workbook

    def crash_staged(path, canonical=None):
        if Path(path).name.startswith(".migration-"):
            raise RuntimeError("simulated crash after backup")
        return original_validate(path, canonical)

    monkeypatch.setattr(data_directory_service, "_validate_workbook", crash_staged)
    with pytest.raises(RuntimeError, match="Could not safely prepare"):
        migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    assert source.is_file()
    assert list((sync_dir / "Backups" / "migration").glob("*/src/*/bank_transactions.xlsx"))

    monkeypatch.setattr(data_directory_service, "_validate_workbook", original_validate)
    report = migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    assert report["complete"] is True
    assert source.is_file()


def test_restart_after_migration_before_cleanup_finishes_only_named_sources(tmp_path, monkeypatch):
    from backend.services.data_migration_service import run_pending_sync_metadata_backfill
    from backend.services.shared_sync_service import synchronize_data_directories

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    _write_root_workbooks(root)
    report = migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    assert all(Path(source).is_file() for source in report["sourceBackups"])
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(sync_dir))
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(sync_dir / "Backups"))
    run_pending_sync_metadata_backfill(data_dir)
    synchronize_data_directories(data_dir, sync_dir)
    report = json.loads((data_dir / ".finledge-data-migration.json").read_text(encoding="utf-8"))

    cleanup = cleanup_migrated_root_sources(data_dir, [root], report, shared_sync_dir=sync_dir)
    assert cleanup["status"] == "complete"
    assert not list(root.glob("*.xlsx"))


def test_cleanup_crash_midway_is_resumed_without_reimporting_deleted_sources(tmp_path, monkeypatch):
    from backend.services.data_migration_service import run_pending_sync_metadata_backfill
    from backend.services.shared_sync_service import synchronize_data_directories
    import pytest

    root = tmp_path / "root"
    data_dir = root / "FinLedge Data"
    sync_dir = root / "FinLedge Data Shared"
    root.mkdir()
    _write_root_workbooks(root)
    report = migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    monkeypatch.setenv("FINLEDGE_SHARED_SYNC_DIR", str(sync_dir))
    monkeypatch.setenv("FINLEDGE_BACKUP_DIR", str(sync_dir / "Backups"))
    run_pending_sync_metadata_backfill(data_dir)
    synchronize_data_directories(data_dir, sync_dir)
    report = json.loads((data_dir / ".finledge-data-migration.json").read_text(encoding="utf-8"))
    second_source = root / "share_transactions.xlsx"
    real_unlink = Path.unlink
    interrupted = {"done": False}

    def unlink_then_crash(path, *args, **kwargs):
        result = real_unlink(path, *args, **kwargs)
        if Path(path) == second_source and not interrupted["done"]:
            interrupted["done"] = True
            raise RuntimeError("simulated process crash during cleanup")
        return result

    monkeypatch.setattr(Path, "unlink", unlink_then_crash)
    with pytest.raises(RuntimeError, match="simulated process crash"):
        cleanup_migrated_root_sources(data_dir, [root], report, shared_sync_dir=sync_dir)
    assert not second_source.exists()
    monkeypatch.setattr(Path, "unlink", real_unlink)

    resumed = cleanup_migrated_root_sources(data_dir, [root], report, shared_sync_dir=sync_dir)
    assert resumed["status"] == "complete"
    assert not list(root.glob("*.xlsx"))
    next_start = migrate_legacy_workbooks(data_dir, [root], sync_dir / "Backups", cleanup_root_dirs=[root], shared_sync_dir=sync_dir)
    assert not any(item.get("source") and Path(item["source"]).parent == root for item in next_start["files"])
    bank = load_workbook(data_dir / "bank_transactions.xlsx", read_only=True, data_only=True)
    try:
        assert bank["Bank"].max_row == 2
    finally:
        bank.close()

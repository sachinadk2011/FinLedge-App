"""Isolated transfer import and balance validation regressions."""

from concurrent.futures import ThreadPoolExecutor
from datetime import date
from threading import Barrier

import pytest
from openpyxl import Workbook, load_workbook

from backend.services import personal_finance_service, personal_finance_transfers
from backend.services.settings_data import config as settings_config
from backend.services.settings_data import workbook_export, workbook_import
from backend.services.transfer_directions import (
    destination_flow,
    normalize_transfer_direction,
    origin_flow,
)


@pytest.fixture
def transfer_data_dir(tmp_path, monkeypatch):
    data_dir = tmp_path / "FinLedge Data"
    data_dir.mkdir()
    bank_path = data_dir / "personal_finance_bank_flow.xlsx"
    cash_path = data_dir / "personal_finance_cash_flow.xlsx"
    transfer_path = data_dir / "personal_finance_transfer.xlsx"

    monkeypatch.setattr(personal_finance_service, "DATA_DIR", data_dir)
    monkeypatch.setattr(personal_finance_service, "BANK_FLOW_FILE_PATH", bank_path)
    monkeypatch.setattr(personal_finance_service, "CASH_FLOW_FILE_PATH", cash_path)
    monkeypatch.setattr(personal_finance_service, "TRANSFER_FILE_PATH", transfer_path)
    monkeypatch.setattr(personal_finance_transfers, "DATA_DIR", data_dir)
    monkeypatch.setattr(personal_finance_transfers, "TRANSFER_FILE_PATH", transfer_path)
    monkeypatch.setattr(workbook_import, "DATA_DIR", data_dir)

    transfer_config = dict(settings_config.DATA_TYPE_CONFIG["pf-transfer"])
    transfer_config["path"] = transfer_path
    monkeypatch.setitem(settings_config.DATA_TYPE_CONFIG, "pf-transfer", transfer_config)
    return data_dir


def _add_income(flow_type: str, amount: float, entry_date: date = date(2025, 1, 1)):
    return personal_finance_service.append_personal_finance_record(
        entry_date=entry_date,
        flow_type=flow_type,
        direction="income",
        category="Salary",
        amount=amount,
        description=f"{flow_type} income",
    )


def _transfer_xlsx(path, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transfer Flow"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _import(path, mode="replace"):
    return workbook_import.import_data_file("pf-transfer", path, mode)


@pytest.mark.parametrize(
    ("alias", "canonical", "origin", "destination"),
    [
        ("Bank to Cash", "bank_to_cash", "bank", "cash"),
        ("bank -> cash", "bank_to_cash", "bank", "cash"),
        ("withdraw", "bank_to_cash", "bank", "cash"),
        ("Cash to Bank", "cash_to_bank", "cash", "bank"),
        ("deposit", "cash_to_bank", "cash", "bank"),
    ],
)
def test_shared_direction_rules_normalize_aliases(alias, canonical, origin, destination):
    assert normalize_transfer_direction(alias) == canonical
    assert origin_flow(alias) == origin
    assert destination_flow(alias) == destination


def test_finledge_transfer_export_round_trips_and_remains_visible(transfer_data_dir):
    _add_income("bank", 500)
    personal_finance_transfers.create_transfer_record(
        date(2025, 2, 1), "Bank to Cash", 125, "export round trip"
    )
    exported = workbook_export.export_data_file("pf-transfer")

    result = _import(exported)
    records = personal_finance_service.read_personal_finance_records("bank")
    transfers = [record for record in records if record["source"] == "transfer"]

    assert result["imported_rows"] == 1
    assert len(transfers) == 1
    assert transfers[0]["transfer_direction"] == "bank_to_cash"
    assert any(
        record.get("source") == "transfer" and record.get("transfer_direction") == "bank_to_cash"
        for record in personal_finance_service.read_personal_finance_records("combined")
    )


def test_transfer_reader_copies_legacy_personal_finance_sheet(transfer_data_dir):
    legacy = Workbook()
    sheet = legacy.active
    sheet.title = "Personal Finance"
    sheet.append(personal_finance_service.HEADERS)
    sheet.append([
        "2025-02-01", "bank_to_cash", "transfer", "Transfer", 25, 25,
        "legacy sheet", "transfer", "2025-02-01 10:00:00", "2025-02-01 10:00:00",
        "transfer:legacy", "desktop",
    ])
    legacy.save(personal_finance_service.TRANSFER_FILE_PATH)
    legacy.close()

    records = personal_finance_transfers.read_transfer_records()

    assert len(records) == 1
    assert records[0]["transfer_direction"] == "bank_to_cash"
    workbook = load_workbook(personal_finance_service.TRANSFER_FILE_PATH, read_only=True)
    assert "Transfer Flow" in workbook.sheetnames
    assert "Personal Finance" in workbook.sheetnames
    workbook.close()
    # Opening through the service a second time must not duplicate the copied row.
    assert len(personal_finance_transfers.read_transfer_records()) == 1


def test_moving_personal_finance_record_keeps_source_when_append_fails(transfer_data_dir, monkeypatch):
    personal_finance_service.append_personal_finance_record(
        date(2025, 1, 1), "bank", "expense", "Food", 20, "move me"
    )
    original = personal_finance_service.read_personal_finance_records("bank")[0]

    def fail_append(**kwargs):
        raise OSError("simulated destination write failure")

    monkeypatch.setattr(personal_finance_service, "append_personal_finance_record", fail_append)
    with pytest.raises(OSError, match="simulated destination write failure"):
        personal_finance_service.update_personal_finance_record(
            original["id"], "bank", date(2025, 1, 2), "cash", "expense", "Food", 30, "move me"
        )

    remaining = personal_finance_service.read_personal_finance_records("bank")
    assert len(remaining) == 1
    assert remaining[0]["description"] == "move me"


def test_moving_personal_finance_record_appends_then_removes_source(transfer_data_dir, monkeypatch):
    personal_finance_service.append_personal_finance_record(
        date(2025, 1, 1), "bank", "expense", "Food", 20, "move me"
    )
    original = personal_finance_service.read_personal_finance_records("bank")[0]
    snapshots = []
    real_snapshot = personal_finance_service.snapshot_before_write

    def record_snapshot(path):
        snapshots.append(path)
        return real_snapshot(path)

    monkeypatch.setattr(personal_finance_service, "snapshot_before_write", record_snapshot)

    updated = personal_finance_service.update_personal_finance_record(
        original["id"], "bank", date(2025, 1, 2), "cash", "expense", "Food", 30, "moved"
    )

    assert personal_finance_service.read_personal_finance_records("bank") == []
    moved = personal_finance_service.read_personal_finance_records("cash")
    assert len(moved) == 1
    assert moved[0]["description"] == "moved"
    assert updated["flow_type"] == "cash"
    assert personal_finance_service.BANK_FLOW_FILE_PATH in snapshots
    assert personal_finance_service.CASH_FLOW_FILE_PATH in snapshots


def test_import_accepts_direction_column_with_canonical_value(tmp_path, transfer_data_dir):
    _add_income("bank", 500)
    source = _transfer_xlsx(
        tmp_path / "direction.xlsx",
        ["Date", "Direction", "Amount"],
        [["2025-02-01", "bank_to_cash", 100]],
    )

    _import(source)

    assert personal_finance_transfers.read_transfer_records()[0]["transfer_direction"] == "bank_to_cash"


def test_import_accepts_file_with_only_transfer_direction_column(tmp_path, transfer_data_dir):
    _add_income("cash", 300)
    source = _transfer_xlsx(
        tmp_path / "transfer-direction.xlsx",
        ["Date", "Transfer Direction", "Amount"],
        [["2025-02-01", "Cash to Bank", 75]],
    )

    _import(source)

    assert personal_finance_transfers.read_transfer_records()[0]["transfer_direction"] == "cash_to_bank"


def test_invalid_transfer_directions_report_row_count(tmp_path, transfer_data_dir):
    source = _transfer_xlsx(
        tmp_path / "invalid.xlsx",
        ["Date", "Direction", "Amount"],
        [
            ["2025-02-01", "sideways", 10],
            ["2025-02-02", "", 20],
            ["2025-02-03", "withdraw", 5],
        ],
    )

    with pytest.raises(ValueError, match="2 rows have a missing or invalid transfer direction"):
        _import(source)


def test_share_workbooks_are_rejected_for_transfer_import(tmp_path, transfer_data_dir):
    source = _transfer_xlsx(
        tmp_path / "shares.xlsx",
        ["Date", "Share Name", "Category", "Per Unit Price", "Allotted", "Buy/Sell"],
        [["2025-02-01", "FinLedge", "buy", 100, 1, "buy"]],
    )

    with pytest.raises(ValueError, match="Share Portfolio workbook"):
        _import(source)


def test_import_rejects_a_transfer_that_overdraws_the_origin(tmp_path, transfer_data_dir):
    _add_income("bank", 100)
    source = _transfer_xlsx(
        tmp_path / "overdraft.xlsx",
        ["Date", "Direction", "Amount"],
        [["2025-02-01", "bank_to_cash", 120]],
    )

    with pytest.raises(ValueError, match=r"overdraw Bank; only NPR 100\.00 is available"):
        _import(source)


def test_import_replays_transfers_chronologically_across_both_flows(tmp_path, transfer_data_dir):
    _add_income("bank", 100)
    source = _transfer_xlsx(
        tmp_path / "chronology.xlsx",
        ["Date", "Direction", "Amount"],
        [
            ["2025-02-01", "bank_to_cash", 80],
            ["2025-02-02", "cash_to_bank", 40],
            ["2025-02-03", "bank_to_cash", 61],
        ],
    )

    with pytest.raises(ValueError, match=r"overdraw Bank; only NPR 60\.00 is available"):
        _import(source)


def test_update_validates_new_amount_against_income_excluding_itself(transfer_data_dir):
    _add_income("bank", 100)
    personal_finance_transfers.create_transfer_record(date(2025, 2, 1), "bank_to_cash", 50)

    with pytest.raises(ValueError, match=r"Only NPR 100\.00 is available in Bank"):
        personal_finance_transfers.update_transfer_record(1, date(2025, 2, 1), "bank_to_cash", 101)


def test_update_direction_change_removes_the_old_receiving_effect(transfer_data_dir):
    _add_income("bank", 100)
    _add_income("cash", 20)
    personal_finance_transfers.create_transfer_record(date(2025, 2, 1), "bank_to_cash", 50)

    with pytest.raises(ValueError, match=r"Only NPR 20\.00 is available in Cash"):
        personal_finance_transfers.update_transfer_record(1, date(2025, 2, 1), "cash_to_bank", 21)


def test_simultaneous_transfers_cannot_spend_the_same_income_twice(transfer_data_dir):
    _add_income("bank", 100)
    barrier = Barrier(2)

    def create_transfer(_):
        barrier.wait()
        try:
            return personal_finance_transfers.create_transfer_record(
                date(2025, 2, 1), "bank_to_cash", 60
            )
        except ValueError as error:
            return error

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(create_transfer, range(2)))

    assert sum(not isinstance(result, ValueError) for result in results) == 1
    assert sum(isinstance(result, ValueError) for result in results) == 1
    assert len(personal_finance_transfers.read_transfer_records()) == 1

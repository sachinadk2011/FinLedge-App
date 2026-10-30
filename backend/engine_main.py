import ctypes
import os
import argparse
import json
import sys
import threading
import time
import traceback
from pathlib import Path

import uvicorn


def _prepare_data_directory() -> dict:
    from backend.services.data_directory_service import migrate_legacy_workbooks
    from backend.services.data_migration_service import record_migration_state, run_pending_data_migrations
    from backend.services.shared_sync_service import synchronize_data_directories
    from backend.services.path_utils import get_data_dir
    from backend.services.workbook_initialization import initialize_data_workbooks
    from backend.services.excel_utils import mirror_data_directory_to_safety_copy, snapshot_before_write
    import shutil
    import uuid
    from openpyxl import load_workbook

    raw_sources = os.getenv("FINLEDGE_LEGACY_DATA_DIRS", "[]")
    try:
        sources = [Path(item) for item in json.loads(raw_sources)]
    except (TypeError, ValueError) as exc:
        raise RuntimeError("Invalid legacy data folder configuration") from exc
    raw_safety_root = str(os.getenv("FINLEDGE_SAFETY_COPY_DIR") or "").strip()
    safety_root = Path(raw_safety_root).expanduser() if raw_safety_root else None
    raw_backup_root = str(os.getenv("FINLEDGE_BACKUP_DIR") or "").strip()
    backup_root = Path(raw_backup_root).expanduser() if raw_backup_root else None
    # Create shared safety/backup locations first. For development and Testing,
    # the parent directory is the already-shared runtime root.
    if safety_root is not None:
        safety_root.mkdir(parents=True, exist_ok=True)
    if backup_root is not None:
        backup_root.mkdir(parents=True, exist_ok=True)
    data_dir = get_data_dir(create=False)
    data_dir.parent.mkdir(parents=True, exist_ok=True)
    try:
        cleanup_roots = [Path(item) for item in json.loads(os.getenv("FINLEDGE_CLEANUP_LEGACY_ROOTS", "[]"))]
    except (TypeError, ValueError) as exc:
        raise RuntimeError("Invalid legacy cleanup folder configuration") from exc
    restored_from_safety = []
    if safety_root is not None and safety_root.is_dir():
        # Recover only a missing or truly empty destination workbook. A populated
        # working copy remains authoritative, so old safety copies cannot revive
        # transactions the user deliberately removed.
        for safety_source in safety_root.glob("*.xlsx"):
            destination = data_dir / safety_source.name
            destination_has_rows = False
            if destination.is_file():
                try:
                    workbook = load_workbook(destination, read_only=True, data_only=True)
                    try:
                        destination_has_rows = any(
                            any(value not in (None, "") for value in row)
                            for sheet in workbook.worksheets
                            for row in sheet.iter_rows(min_row=2, values_only=True)
                        )
                    finally:
                        workbook.close()
                except Exception:
                    destination_has_rows = True
            if not destination_has_rows:
                data_dir.mkdir(parents=True, exist_ok=True)
                if destination.is_file():
                    snapshot_before_write(destination)
                staging = data_dir / f".{safety_source.stem}.{uuid.uuid4().hex}.restore.xlsx"
                try:
                    shutil.copy2(safety_source, staging)
                    try:
                        check = load_workbook(staging, read_only=True, data_only=True)
                    except Exception:
                        continue
                    else:
                        check.close()
                    os.replace(staging, destination)
                    restored_from_safety.append(safety_source.name)
                finally:
                    staging.unlink(missing_ok=True)
    migration = migrate_legacy_workbooks(
        data_dir,
        sources,
        backup_root,
        cleanup_root_dirs=cleanup_roots,
        shared_sync_dir=Path(os.getenv("FINLEDGE_SHARED_SYNC_DIR")) if os.getenv("FINLEDGE_SHARED_SYNC_DIR") else None,
    )
    schema_migration = run_pending_data_migrations(data_dir)
    workbooks = initialize_data_workbooks(data_dir)
    shared_sync_dir = Path(os.getenv("FINLEDGE_SHARED_SYNC_DIR")) if os.getenv("FINLEDGE_SHARED_SYNC_DIR") else None
    sync_report = synchronize_data_directories(data_dir, shared_sync_dir)
    safety_copies = mirror_data_directory_to_safety_copy(data_dir)
    cleanup_report = {"status": "deferred_until_metadata_backfill", "roots": [str(path) for path in cleanup_roots]}
    # Initialization may create missing workbooks. Record their final fingerprints
    # so FastAPI's lifespan check can return without opening the files again.
    record_migration_state(data_dir)
    return {
        **migration,
        "schemaMigration": schema_migration,
        "workbooks": workbooks,
        "restoredFromSafetyCopy": restored_from_safety,
        "safetyCopies": safety_copies,
        "legacyCleanup": cleanup_report,
        "sharedSync": sync_report,
    }


def _watch_parent_process() -> None:
    """Exit the sidecar if its Electron owner is gone (including force-close)."""
    raw_parent_pid = str(os.getenv("FINLEDGE_PARENT_PID") or "").strip()
    if not raw_parent_pid.isdigit():
        return
    parent_pid = int(raw_parent_pid)

    def monitor() -> None:
        if os.name == "nt":
            kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
            kernel32.OpenProcess.argtypes = (ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32)
            kernel32.OpenProcess.restype = ctypes.c_void_p
            kernel32.WaitForSingleObject.argtypes = (ctypes.c_void_p, ctypes.c_uint32)
            kernel32.WaitForSingleObject.restype = ctypes.c_uint32
            kernel32.CloseHandle.argtypes = (ctypes.c_void_p,)
            kernel32.CloseHandle.restype = ctypes.c_int
            synchronize = 0x00100000
            wait_object_0 = 0
            wait_timeout = 258

            while True:
                handle = kernel32.OpenProcess(synchronize, False, parent_pid)
                if not handle:
                    # Windows reports ERROR_INVALID_PARAMETER when the PID no longer exists.
                    if ctypes.get_last_error() == 87:
                        os._exit(0)
                    time.sleep(2)
                    continue
                try:
                    wait_result = kernel32.WaitForSingleObject(handle, 0)
                finally:
                    kernel32.CloseHandle(handle)
                if wait_result == wait_object_0:
                    os._exit(0)
                if wait_result != wait_timeout:
                    time.sleep(2)
                    continue
                time.sleep(2)

        while True:
            try:
                os.kill(parent_pid, 0)
            except OSError:
                os._exit(0)
            time.sleep(2)

    threading.Thread(target=monitor, name="finledge-parent-watch", daemon=True).start()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.description = "Start the FinLedge API, or prepare data for isolated tests."
    parser.add_argument("--prepare-data", action="store_true")
    options = parser.parse_args()
    if options.prepare_data:
        print(json.dumps(_prepare_data_directory()))
        return

    _watch_parent_process()
    preparation_started = time.perf_counter()
    try:
        preparation_report = _prepare_data_directory()
    except Exception:
        traceback.print_exc(file=sys.stderr)
        # Electron uses this status to preserve the existing legacy-folder fallback.
        raise SystemExit(75)
    preparation_ms = round((time.perf_counter() - preparation_started) * 1000)
    print(f"FINLEDGE_MIGRATION_REPORT={json.dumps(preparation_report, separators=(',', ':'))}", flush=True)
    print(f"FINLEDGE_STARTUP_PHASE=data-ready elapsed_ms={preparation_ms}", flush=True)

    from backend.main import app

    host = os.getenv("FINLEDGE_BACKEND_HOST", "127.0.0.1")
    port = int(os.getenv("FINLEDGE_BACKEND_PORT", "8000"))
    log_level = os.getenv("FINLEDGE_BACKEND_LOG_LEVEL", "info")

    uvicorn.run(app, host=host, port=port, log_level=log_level)


if __name__ == "__main__":
    main()
